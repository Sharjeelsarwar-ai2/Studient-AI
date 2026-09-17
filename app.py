# Build the user's full app with ONLY the flashcard-related changes applied.
# This keeps the rest of the supplied code unchanged while making the flashcards
# generate visual artifacts and connecting the Flashcards tab to generate_flashcards().


""" 
Studient AI 
================= 
Flat single-row tab layout (matches the original UI), with the graded 
interactive Test and Analytics features folded in as regular tabs. 
 
Requires: streamlit>=1.45.0, groq, pypdf, python-docx, python-pptx 
""" 
 
import streamlit as st 
from pypdf import PdfReader 
from groq import Groq 

# Optional readers for Word and PowerPoint uploads. The app remains import-safe
# when a deployment has not installed these packages yet.
try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

import re 
import json 
import os 
import hashlib
import random
import time
import zipfile
import textwrap
import base64
import secrets
import html as _html_escape_lib  # aliased: this file defines its own html() render helper below 
from io import BytesIO, StringIO
import csv
from datetime import datetime, date, timedelta
from collections import Counter, defaultdict 

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.units import inch
except ImportError:
    SimpleDocTemplate = None

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:
    Image = ImageDraw = ImageFont = None

try:
    from gtts import gTTS
except ImportError:
    gTTS = None

try:
    import pytesseract
    from pdf2image import convert_from_bytes
except ImportError:
    pytesseract = None
    convert_from_bytes = None

try:
    from streamlit_autorefresh import st_autorefresh
except ImportError:
    st_autorefresh = None

try:
    from cryptography.fernet import Fernet, InvalidToken
except ImportError:
    Fernet = None
    InvalidToken = Exception
 
# ============================================================ 
# PAGE CONFIG 
# ============================================================ 
 
st.set_page_config( 
    page_title="Studient AI", 
    page_icon="🧠", 
    layout="wide", 
    initial_sidebar_state="expanded", 
) 
 
# ============================================================ 
# HTML HELPER  (avoids the markdown-code-block bug entirely) 
# ============================================================ 
 
def html(content: str): 
    if hasattr(st, "html"): 
        st.html(content) 
    else: 
        st.markdown("\n".join(line.lstrip() for line in content.splitlines()), 
                     unsafe_allow_html=True) 
 
 
# ============================================================ 
# CSS  -- richer color, left-aligned to dodge markdown code fences 
# ============================================================ 
 
html(""" 
<style> 
:root { 
  --bg: #eef0fb; 
  --ink: #10142a; 
  --ink-soft: #454e68; 
  --ink-faint: #7d84a0; 
  --accent1: #4f46e5; 
  --accent2: #7c3aed; 
  --accent3: #0ea5e9; 
  --accent4: #db2777; 
  --accent5: #f59e0b; 
  --gold: #d6a92c; 
  --good: #16a34a; 
  --bad: #dc2626; 
  --glass-bg: rgba(255,255,255,0.58); 
  --glass-bg-soft: rgba(255,255,255,0.44); 
  --glass-border: rgba(255,255,255,0.75); 
  --glass-shadow: 0 10px 40px rgba(31,25,90,0.14), inset 0 1px 0 rgba(255,255,255,0.65); 
} 
 
/* PREMIUM MESH BACKGROUND */ 
.stApp { 
  background: 
    radial-gradient(circle at 8% 0%, rgba(99,102,241,0.28), transparent 34%), 
    radial-gradient(circle at 92% 4%, rgba(219,39,119,0.20), transparent 32%), 
    radial-gradient(circle at 50% 30%, rgba(14,165,233,0.16), transparent 40%), 
    radial-gradient(circle at 20% 95%, rgba(214,169,44,0.16), transparent 32%), 
    radial-gradient(circle at 85% 90%, rgba(124,58,237,0.18), transparent 34%), 
    linear-gradient(160deg, #eef0fb 0%, #eaeefc 45%, #f3ecfb 100%); 
  background-attachment: fixed; 
} 
.main .block-container { max-width: 1450px; padding-top: 1.6rem; padding-bottom: 3rem; } 
 
/* HERO -- deep glass with a gold shimmer edge */ 
.sm-hero { 
  position: relative; overflow: hidden; 
  padding: 48px 30px; margin-bottom: 26px; border-radius: 30px; text-align: center; 
  background: linear-gradient(135deg, rgba(255,255,255,0.75), rgba(238,242,255,0.55) 60%, rgba(253,242,255,0.6)); 
  border: 1px solid var(--glass-border); 
  box-shadow: var(--glass-shadow); 
  backdrop-filter: blur(26px) saturate(170%); 
  -webkit-backdrop-filter: blur(26px) saturate(170%); 
} 
.sm-hero::before { 
  content: ""; position: absolute; inset: 0; border-radius: 30px; padding: 1px; 
  background: linear-gradient(120deg, rgba(214,169,44,0.55), rgba(124,58,237,0.25), rgba(14,165,233,0.35)); 
  -webkit-mask: linear-gradient(#fff 0 0) content-box, linear-gradient(#fff 0 0); 
  -webkit-mask-composite: xor; mask-composite: exclude; pointer-events: none; 
} 
.sm-hero h1 { margin: 0; font-size: 47px; font-weight: 850; color: var(--ink); letter-spacing: -1.5px; } 
.sm-hero .grad { 
  background: linear-gradient(90deg, var(--accent1), var(--accent2), var(--accent4), var(--gold)); 
  -webkit-background-clip: text; background-clip: text; -webkit-text-fill-color: transparent; 
} 
.sm-hero p.sub { color: var(--ink-soft); font-size: 18px; font-weight: 650; margin-top: 8px; } 
.sm-hero p.desc { color: var(--ink-faint); max-width: 700px; margin: 10px auto 0; font-size: 15px; } 
 
/* GLASS CARD BASE -- reused everywhere below */ 
.sm-card, .sm-feature, .sm-metric, .sm-question, .sm-answer, .sm-warning, .sm-score-hero { 
  backdrop-filter: blur(20px) saturate(160%); 
  -webkit-backdrop-filter: blur(20px) saturate(160%); 
} 
 
.sm-card { 
  padding: 24px; border-radius: 20px; background: var(--glass-bg); 
  border: 1px solid var(--glass-border); box-shadow: var(--glass-shadow); 
  margin-bottom: 16px; color: var(--ink); 
} 
.sm-card h2, .sm-card h3, .sm-card p, .sm-card li, .sm-card b { color: var(--ink); } 
.sm-card p { color: var(--ink-soft); } 
 
.sm-feature { 
  padding: 22px; border-radius: 20px; background: var(--glass-bg); 
  border: 1px solid var(--glass-border); min-height: 155px; 
  box-shadow: var(--glass-shadow); 
  transition: transform .2s ease; 
} 
.sm-feature:hover { transform: translateY(-3px); } 
.sm-feature .icon { font-size: 30px; margin-bottom: 8px; } 
.sm-feature .title { font-weight: 750; color: var(--ink); font-size: 18px; margin-bottom: 6px; } 
.sm-feature .text { color: var(--ink-soft); font-size: 14px; line-height: 1.55; } 
 
/* METRIC CHIPS -- colored icon badge + big number, premium glass finish */ 
.sm-metric { 
  padding: 20px 22px; border-radius: 20px; background: var(--glass-bg); 
  border: 1px solid var(--glass-border); box-shadow: var(--glass-shadow); 
} 
.sm-metric .row { display: flex; align-items: center; gap: 8px; margin-bottom: 10px; } 
.sm-metric .chip { 
  width: 28px; height: 28px; border-radius: 9px; display: flex; align-items: center; 
  justify-content: center; font-size: 14px; font-weight: 800; 
  box-shadow: inset 0 1px 0 rgba(255,255,255,0.6); 
} 
.sm-metric .label { color: var(--ink-soft); font-weight: 650; font-size: 15px; } 
.sm-metric .value { color: var(--ink); font-weight: 850; font-size: 32px; } 
.chip-blue   { background: linear-gradient(135deg, #c7d2fe, #a5b4fc); color: #3730a3; } 
.chip-cyan   { background: linear-gradient(135deg, #bae6fd, #7dd3fc); color: #075985; } 
.chip-orange { background: linear-gradient(135deg, #fed7aa, #fdba74); color: #9a3412; } 
 
/* QUESTION / ANSWER / WARNING -- glass with colored left rail */ 
.sm-question { 
  padding: 20px; margin-bottom: 14px; border-radius: 18px; background: var(--glass-bg-soft); 
  border: 1px solid rgba(79,70,229,0.18); border-left: 5px solid var(--accent1); 
  box-shadow: var(--glass-shadow); 
} 
.sm-question .label { color: var(--accent1); font-size: 12px; font-weight: 800; text-transform: uppercase; letter-spacing: .6px; margin-bottom: 6px; } 
 
.sm-answer { 
  padding: 22px; margin-top: 14px; border-radius: 18px; background: rgba(220,252,231,0.55); 
  border: 1px solid rgba(34,197,94,0.25); border-left: 5px solid var(--good); 
  box-shadow: var(--glass-shadow); 
} 
.sm-answer .title { color: #14532d; font-weight: 750; font-size: 17px; margin-bottom: 10px; } 
 
.sm-warning { 
  padding: 16px; border-radius: 16px; background: rgba(255,237,213,0.6); 
  border: 1px solid rgba(249,115,22,0.28); color: #7c2d12; 
  box-shadow: var(--glass-shadow); 
} 
 
/* QUIZ */ 
.sm-quiz-progress { color: var(--ink-faint); font-size: 13px; font-weight: 750; text-transform: uppercase; letter-spacing: .5px; } 
.sm-quiz-topic { 
  display: inline-block; padding: 4px 13px; border-radius: 999px; font-size: 12px; font-weight: 750; 
  background: linear-gradient(135deg, rgba(219,39,119,0.16), rgba(124,58,237,0.14)); 
  color: var(--accent4); margin-bottom: 10px; border: 1px solid rgba(219,39,119,0.2); 
} 
.sm-quiz-q { font-size: 20px; font-weight: 750; color: var(--ink); margin-bottom: 6px; } 
.sm-result-correct { color: var(--good); font-weight: 750; } 
.sm-result-wrong { color: var(--bad); font-weight: 750; } 
 
.sm-score-hero { 
  text-align: center; padding: 40px; border-radius: 26px; 
  background: linear-gradient(135deg, rgba(220,252,231,0.55), rgba(238,242,255,0.55), rgba(253,242,255,0.55)); 
  border: 1px solid var(--glass-border); box-shadow: var(--glass-shadow); 
} 
.sm-score-hero .big { 
  font-size: 60px; font-weight: 850; 
  background: linear-gradient(90deg, var(--accent1), var(--gold)); 
  -webkit-background-clip: text; background-clip: text; -webkit-text-fill-color: transparent; 
} 
.sm-score-hero .pct { font-size: 21px; font-weight: 750; color: var(--accent1); } 
 
.sm-tag-weak { background: rgba(220,38,38,0.12); color: var(--bad); padding: 4px 12px; border-radius: 999px; font-size: 13px; font-weight: 750; margin: 3px; display: inline-block; border: 1px solid rgba(220,38,38,0.2); } 
.sm-tag-strong { background: rgba(22,163,74,0.12); color: var(--good); padding: 4px 12px; border-radius: 999px; font-size: 13px; font-weight: 750; margin: 3px; display: inline-block; border: 1px solid rgba(22,163,74,0.2); } 
 
/* BUTTONS -- premium gradient with gold-tinted glow */ 
.stButton > button { 
  width: 100%; min-height: 46px; border: none; border-radius: 13px; font-weight: 750; color: white !important; 
  background: linear-gradient(135deg, var(--accent1) 0%, var(--accent2) 55%, var(--accent4) 100%); 
  box-shadow: 0 10px 24px rgba(124,58,237,0.30), 0 0 0 1px rgba(255,255,255,0.15) inset; 
  transition: transform .15s ease, box-shadow .15s ease; 
} 
.stButton > button:hover { transform: translateY(-2px); box-shadow: 0 16px 32px rgba(124,58,237,0.38), 0 0 0 1px rgba(214,169,44,0.4) inset; } 
 
/* SIDEBAR -- frosted glass panel */ 
section[data-testid="stSidebar"] { 
  background: linear-gradient(180deg, rgba(255,255,255,0.75), rgba(238,240,251,0.7)); 
  backdrop-filter: blur(18px); -webkit-backdrop-filter: blur(18px); 
  border-right: 1px solid rgba(255,255,255,0.5); 
} 
[data-testid="stFileUploader"] { 
  padding: 8px; border-radius: 18px; background: rgba(255,255,255,0.5); 
  border: 1.5px dashed rgba(124,58,237,0.4); backdrop-filter: blur(10px); 
} 
[data-testid="stMetric"] { 
  padding: 14px; border-radius: 16px; background: rgba(255,255,255,0.6); 
  border: 1px solid var(--glass-border); backdrop-filter: blur(14px); 
} 
 
/* FLAT SCROLLABLE TAB ROW -- glass pill bar */ 
.stTabs [data-baseweb="tab-list"] { 
  gap: 4px; padding: 7px; border-radius: 18px; background: var(--glass-bg); 
  border: 1px solid var(--glass-border); box-shadow: var(--glass-shadow); 
  backdrop-filter: blur(20px); -webkit-backdrop-filter: blur(20px); 
  overflow-x: auto; flex-wrap: nowrap; 
} 
.stTabs [data-baseweb="tab"] { border-radius: 999px; padding: 10px 18px; font-weight: 600; border:none; color: var(--ink-soft); white-space: nowrap; } 
.stTabs [aria-selected="true"] { 
  color: white !important; 
  background: linear-gradient(135deg, var(--accent1), var(--accent2)) !important;
  border-radius:999px !important;
  padding: 10px 18px !important;
  border: none !important;
  box-shadow: 0 4px 12px rgba(0,0,0,0,12) !important;
  
} 
 
/* FLASHCARDS -- real flip cards, pure CSS (checkbox hack), no JS */ 
/* FLASHCARDS -- interactive visual flip cards */
.sm-flip-grid {
  display: grid;
  grid-template-columns: repeat(auto-fill, minmax(300px, 1fr));
  gap: 22px;
  margin-top: 15px;
}

.sm-flip-card {
  position: relative;
  height: 330px;
  perspective: 1400px;
  animation: smCardRise .5s ease both;
}

.sm-flip-card:nth-child(2) { animation-delay: .06s; }
.sm-flip-card:nth-child(3) { animation-delay: .12s; }
.sm-flip-card:nth-child(4) { animation-delay: .18s; }
.sm-flip-card:nth-child(5) { animation-delay: .24s; }
@keyframes smCardRise { from { opacity:0; transform:translateY(10px) scale(.985); } to { opacity:1; transform:translateY(0) scale(1); } }

.sm-flip-card input[type="checkbox"] {
  display: none;
}

.sm-flip-inner {
  display: block;
  position: relative;
  width: 100%;
  height: 100%;
  cursor: pointer;
  transform-style: preserve-3d;
  transition: transform 0.65s cubic-bezier(.4,.2,.2,1);
}

.sm-flip-card input:checked + .sm-flip-inner {
  transform: rotateY(180deg);
}

.sm-flip-front,
.sm-flip-back {
  position: absolute;
  inset: 0;
  backface-visibility: hidden;
  -webkit-backface-visibility: hidden;
  border-radius: 22px;
  padding: 22px;
  display: flex;
  flex-direction: column;
  border: 1px solid var(--glass-border);
  box-shadow: var(--glass-shadow);
  backdrop-filter: blur(20px) saturate(160%);
  -webkit-backdrop-filter: blur(20px) saturate(160%);
  overflow: hidden;
}

.sm-flip-front {
  background:
    radial-gradient(circle at 92% 8%, rgba(219,39,119,0.27), transparent 29%),
    radial-gradient(circle at 4% 96%, rgba(14,165,233,0.20), transparent 32%),
    linear-gradient(145deg, rgba(224,231,255,0.86), rgba(255,255,255,0.58) 56%, rgba(253,242,255,0.74));
  justify-content: space-between;
}

.sm-flip-back {
  background:
    radial-gradient(circle at 8% 8%, rgba(14,165,233,0.28), transparent 30%),
    radial-gradient(circle at 94% 94%, rgba(22,163,74,0.18), transparent 30%),
    linear-gradient(145deg, rgba(220,252,231,0.82), rgba(255,255,255,0.62) 54%, rgba(224,231,255,0.72));
  transform: rotateY(180deg);
  justify-content: space-between;
}

.sm-flip-front::before, .sm-flip-back::before { content:""; position:absolute; left:0; right:0; top:0; height:5px; background:linear-gradient(90deg,var(--accent3),var(--accent1),var(--accent4),var(--gold)); opacity:.9; }
.sm-flip-card:hover .sm-flip-front, .sm-flip-card:hover .sm-flip-back { box-shadow:0 20px 46px rgba(31,25,90,.2), inset 0 1px 0 rgba(255,255,255,.85); }
.sm-flip-card:focus-within { outline:3px solid rgba(14,165,233,.22); outline-offset:5px; border-radius:24px; }

.sm-flip-q {
  font-size: 20px;
  font-weight: 800;
  color: var(--ink);
  line-height: 1.4;
  margin-top: 12px;
}

.sm-flip-hint {
  font-size: 12px;
  color: var(--ink-faint);
  font-weight: 650;
  text-align: center;
}

/* VISUAL ARTIFACT AREA */
.sm-visual-artifact {
  flex: 1;
  display: flex;
  flex-direction: column;
  justify-content: center;
  gap: 12px;
  margin: 8px 0;
}

.sm-visual-title {
  text-align: center;
  font-size: 16px;
  font-weight: 850;
  color: var(--accent1);
  margin-bottom: 4px;
}

.sm-visual-definition {
  display: flex;
  align-items: center;
  justify-content: center;
  min-height: 95px;
}

.sm-visual-core {
  padding: 18px 22px;
  border-radius: 18px;
  background: linear-gradient(135deg, rgba(79,70,229,0.16), rgba(124,58,237,0.18));
  border: 2px solid rgba(79,70,229,0.25);
  font-size: 19px;
  font-weight: 850;
  text-align: center;
  color: var(--ink);
  box-shadow: 0 8px 25px rgba(79,70,229,0.12);
}

.sm-visual-items {
  display: flex;
  flex-direction: column;
  gap: 7px;
}

.sm-visual-item {
  padding: 9px 12px;
  border-radius: 12px;
  background: rgba(255,255,255,0.58);
  border: 1px solid rgba(79,70,229,0.14);
  font-size: 12px;
  font-weight: 650;
  color: var(--ink-soft);
}

.sm-process {
  display: flex;
  flex-direction: column;
  align-items: center;
  gap: 5px;
}

.sm-process-step {
  width: 90%;
  padding: 9px 10px;
  border-radius: 12px;
  background: rgba(255,255,255,0.65);
  border: 1px solid rgba(14,165,233,0.20);
  text-align: center;
  font-size: 12px;
  font-weight: 750;
  color: var(--ink);
}

.sm-arrow {
  font-size: 17px;
  font-weight: 900;
  color: var(--accent2);
  line-height: 1;
}

.sm-comparison {
  display: grid;
  grid-template-columns: 1fr 1fr;
  gap: 8px;
}

.sm-comparison-column {
  padding: 10px;
  border-radius: 14px;
  background: rgba(255,255,255,0.60);
  border: 1px solid rgba(124,58,237,0.15);
}

.sm-comparison-title {
  font-size: 12px;
  font-weight: 850;
  color: var(--accent2);
  text-align: center;
  margin-bottom: 7px;
}

.sm-comparison-item {
  font-size: 11px;
  color: var(--ink-soft);
  margin: 5px 0;
  line-height: 1.3;
}

.sm-timeline {
  display: flex;
  flex-direction: column;
  gap: 8px;
}

.sm-timeline-item {
  display: flex;
  align-items: center;
  gap: 9px;
  padding: 7px 10px;
  border-radius: 11px;
  background: rgba(255,255,255,0.58);
}

.sm-timeline-number {
  min-width: 25px;
  height: 25px;
  border-radius: 50%;
  display: flex;
  align-items: center;
  justify-content: center;
  background: linear-gradient(135deg, var(--accent1), var(--accent2));
  color: white;
  font-size: 11px;
  font-weight: 850;
}

.sm-timeline-text {
  font-size: 11px;
  font-weight: 650;
  color: var(--ink-soft);
}

.sm-formula {
  display: flex;
  align-items: center;
  justify-content: center;
  min-height: 90px;
  padding: 15px;
  border-radius: 17px;
  background: rgba(255,255,255,0.65);
  border: 2px dashed rgba(214,169,44,0.45);
  font-size: 21px;
  font-weight: 850;
  text-align: center;
  color: var(--ink);
}

.sm-fact {
  display: flex;
  align-items: center;
  justify-content: center;
  min-height: 90px;
  padding: 18px;
  border-radius: 18px;
  background: linear-gradient(135deg, rgba(245,158,11,0.18), rgba(219,39,119,0.12));
  border: 1px solid rgba(245,158,11,0.28);
  text-align: center;
  font-size: 17px;
  font-weight: 800;
  color: var(--ink);
}

.sm-flip-a-label {
  font-size: 11px;
  font-weight: 800;
  text-transform: uppercase;
  letter-spacing: .6px;
  color: var(--gold);
  text-align: center;
}

.sm-flip-back-note {
  font-size: 10px;
  color: var(--ink-faint);
  text-align: center;
  line-height: 1.3;
}

@media (max-width: 768px) {
  .sm-hero {
    padding: 32px 16px;
  }

  .sm-hero h1 {
    font-size: 32px;
  }

  .sm-flip-card {
    height: 340px;
  }
}

/* BRAND / NAVIGATION POLISH */
.sm-brand { display:flex; align-items:center; gap:12px; margin:4px 0 22px; }
.sm-brand-mark { width:44px; height:44px; display:grid; place-items:center; border-radius:14px; color:white; font-size:23px; background:linear-gradient(135deg,#4f46e5,#7c3aed 58%,#db2777); box-shadow:0 10px 24px rgba(79,70,229,.28); }
.sm-brand-name { font-weight:850; letter-spacing:-.5px; color:var(--ink); font-size:22px; line-height:1; }
.sm-brand-name span { background:linear-gradient(90deg,var(--accent1),var(--accent4)); -webkit-background-clip:text; background-clip:text; -webkit-text-fill-color:transparent; }
.sm-brand-sub { display:block; margin-top:5px; color:var(--ink-faint); font-size:11px; font-weight:700; letter-spacing:.4px; text-transform:uppercase; }
.sm-hero .eyebrow { display:inline-flex; align-items:center; gap:7px; padding:7px 13px; border-radius:999px; background:rgba(79,70,229,.09); color:var(--accent1); font-size:12px; font-weight:800; letter-spacing:.5px; text-transform:uppercase; }
.sm-hero .hero-icon { display:inline-grid; place-items:center; width:70px; height:70px; margin:18px auto 14px; border-radius:22px; font-size:34px; background:linear-gradient(135deg,rgba(79,70,229,.14),rgba(219,39,119,.12)); border:1px solid rgba(124,58,237,.16); }
.sm-upload-head { margin:2px 0 4px; color:var(--ink); font-size:19px; font-weight:850; }
.sm-upload-copy { color:var(--ink-faint); font-size:12px; line-height:1.5; margin-bottom:12px; }
.sm-file-types { display:flex; flex-wrap:wrap; gap:5px; margin:8px 0 12px; }
.sm-file-type { padding:4px 8px; border-radius:7px; background:rgba(79,70,229,.08); border:1px solid rgba(79,70,229,.12); color:var(--accent1); font-size:10px; font-weight:800; }
[data-testid="stFileUploader"] section { border-radius:15px; border-color:rgba(124,58,237,.35); background:rgba(255,255,255,.42); }
[data-testid="stFileUploader"] small { color:var(--ink-faint); }
.stDownloadButton > button { border-radius:13px; }
.stTextInput input, .stSelectbox [data-baseweb="select"], .stNumberInput input { border-radius:12px; }
@media (max-width: 768px) { .sm-brand { margin-bottom:14px; } .sm-hero .hero-icon { width:58px; height:58px; font-size:28px; } }

/* INTERACTIVE MIND MAP */
.sm-mindmap-shell { margin: 18px 0 6px; padding: 24px; border-radius: 24px; background: linear-gradient(145deg,rgba(255,255,255,.68),rgba(238,242,255,.55)); border: 1px solid var(--glass-border); box-shadow: var(--glass-shadow); overflow-x:auto; }
.sm-mindmap-toolbar { display:flex; align-items:center; justify-content:space-between; gap:12px; flex-wrap:wrap; margin-bottom:18px; }
.sm-mindmap-note { color:var(--ink-faint); font-size:12px; font-weight:650; }
.sm-mindmap-tree, .sm-mindmap-tree ul { list-style:none; margin:0; padding:0; }
.sm-mindmap-tree ul { position:relative; margin-left:30px; padding-left:26px; border-left:2px solid rgba(124,58,237,.18); }
.sm-mindmap-tree li { position:relative; margin:13px 0; min-width:210px; }
.sm-mindmap-tree ul > li::before { content:""; position:absolute; left:-27px; top:25px; width:26px; border-top:2px solid rgba(124,58,237,.18); }
.sm-mindmap-node { display:inline-flex; align-items:center; gap:9px; max-width:360px; padding:12px 16px; border-radius:15px; background:rgba(255,255,255,.78); border:1px solid rgba(124,58,237,.16); box-shadow:0 6px 18px rgba(31,25,90,.08); color:var(--ink); font-size:14px; font-weight:700; line-height:1.35; cursor:pointer; transition:transform .18s ease, box-shadow .18s ease, border-color .18s ease; }
.sm-mindmap-details summary::-webkit-details-marker { display:none; }
.sm-mindmap-details[open] > summary .sm-mindmap-toggle { font-size:0; }
.sm-mindmap-details[open] > summary .sm-mindmap-toggle::after { content:"−"; font-size:13px; }
.sm-mindmap-details:not([open]) > summary .sm-mindmap-toggle { font-size:0; }
.sm-mindmap-details:not([open]) > summary .sm-mindmap-toggle::after { content:"+"; font-size:13px; }
.sm-mindmap-node:hover { transform:translateY(-2px); box-shadow:0 10px 24px rgba(79,70,229,.16); border-color:rgba(79,70,229,.38); }
.sm-mindmap-root > .sm-mindmap-node { padding:17px 22px; border:0; color:#fff; font-size:18px; background:linear-gradient(135deg,var(--accent1),var(--accent2) 58%,var(--accent4)); box-shadow:0 12px 28px rgba(79,70,229,.28); }
.sm-mindmap-node .sm-mindmap-dot { flex:0 0 auto; width:9px; height:9px; border-radius:50%; background:linear-gradient(135deg,var(--accent3),var(--accent2)); }
.sm-mindmap-root > .sm-mindmap-node .sm-mindmap-dot { background:rgba(255,255,255,.82); }
.sm-mindmap-children { margin-top:10px !important; }
.sm-mindmap-children.is-collapsed { display:none; }
.sm-mindmap-toggle { display:inline-grid; place-items:center; width:21px; height:21px; border-radius:7px; background:rgba(79,70,229,.1); color:var(--accent1); font-size:13px; font-weight:900; }
.sm-mindmap-root > .sm-mindmap-children { margin-left:36px; }

/* SECOND UI POLISH PASS — preserve Streamlit's classic typography */
.stApp { min-height: 100vh; background-color: #eef0fb; }
.stApp > header, [data-testid="stHeader"] { background: linear-gradient(105deg, rgba(224,231,255,.34), rgba(255,255,255,.18) 42%, rgba(253,242,255,.30) 76%, rgba(254,243,199,.20)) !important; backdrop-filter: blur(38px) saturate(200%); -webkit-backdrop-filter: blur(38px) saturate(200%); border-bottom: 1px solid rgba(255,255,255,.48); box-shadow: 0 8px 30px rgba(31,25,90,.09), inset 0 1px 0 rgba(255,255,255,.62); }
[data-testid="stToolbar"] { background: linear-gradient(135deg, rgba(255,255,255,.42), rgba(199,210,254,.25)); border: 1px solid rgba(255,255,255,.55); border-radius: 14px; backdrop-filter: blur(16px); }
[data-testid="stDecoration"] { background: linear-gradient(90deg, var(--accent3), var(--accent1), var(--accent2), var(--accent4), var(--gold)) !important; height: 4px !important; box-shadow: 0 0 18px rgba(124,58,237,.38); }
[data-testid="stStatusWidget"] { background: linear-gradient(135deg, rgba(255,255,255,.62), rgba(224,231,255,.35)); border: 1px solid rgba(255,255,255,.78); border-radius: 14px; backdrop-filter: blur(18px); }
.main .block-container { max-width: 1500px; padding: 2.25rem 2.25rem 4rem; }
section[data-testid="stSidebar"] > div { padding: 1.5rem 1.15rem 2rem; }
section[data-testid="stSidebar"] [data-testid="stVerticalBlock"] { gap: .7rem; }
.sm-hero { padding: 54px 34px 48px; margin-bottom: 30px; border-radius: 32px; background: linear-gradient(135deg, rgba(255,255,255,.86), rgba(226,232,255,.62) 52%, rgba(253,232,246,.66)); box-shadow: 0 18px 55px rgba(31,25,90,.18), inset 0 1px 0 rgba(255,255,255,.85); }
.sm-hero h1 { font-size: clamp(38px, 4vw, 58px); letter-spacing: -2.4px; }
.sm-hero p.sub { font-size: 19px; letter-spacing: -.2px; }
.sm-hero p.desc { font-size: 15.5px; line-height: 1.7; }
.sm-card { padding: 28px; border-radius: 24px; background: linear-gradient(145deg, rgba(255,255,255,.78), rgba(239,242,255,.56)); box-shadow: 0 15px 38px rgba(31,25,90,.13), inset 0 1px 0 rgba(255,255,255,.8); }
.sm-feature { min-height: 175px; padding: 25px; border-radius: 22px; background: linear-gradient(145deg, rgba(255,255,255,.78), rgba(235,239,255,.55)); box-shadow: 0 13px 30px rgba(31,25,90,.12), inset 0 1px 0 rgba(255,255,255,.82); }
.sm-feature .icon { width: 48px; height: 48px; display: grid; place-items: center; border-radius: 15px; background: rgba(79,70,229,.1); font-size: 25px; }
.sm-feature .title { margin-top: 14px; font-size: 19px; }
.sm-hero, .sm-card, .sm-feature, .sm-metric, .sm-document-bar { animation: smFadeUp .48s ease both; }
@keyframes smFadeUp { from { opacity:0; transform:translateY(7px); } to { opacity:1; transform:translateY(0); } }
.sm-hero::after { content:""; position:absolute; width:180px; height:180px; right:-55px; top:-75px; border-radius:50%; background:rgba(255,255,255,.22); filter:blur(2px); animation:smGlow 7s ease-in-out infinite alternate; pointer-events:none; }
@keyframes smGlow { from { transform:translate3d(-8px,8px,0); opacity:.35; } to { transform:translate3d(12px,-5px,0); opacity:.7; } }
.sm-metric { padding: 22px 24px; border-radius: 22px; background: linear-gradient(145deg, rgba(255,255,255,.82), rgba(232,238,255,.58)); box-shadow: 0 13px 32px rgba(31,25,90,.13), inset 0 1px 0 rgba(255,255,255,.86); transition: transform .18s ease, box-shadow .18s ease; }
.sm-metric:hover { transform: translateY(-2px); box-shadow: 0 16px 34px rgba(31,25,90,.16); }
.sm-metric .value { font-size: clamp(25px, 2.3vw, 35px); letter-spacing: -.8px; }
.stTabs { margin-top: 8px; }
.stTabs [data-baseweb="tab-list"] { gap: 6px; padding: 8px; border-radius: 20px; position: sticky; top: 3.25rem; z-index: 5; background: rgba(255,255,255,.62); box-shadow: 0 12px 30px rgba(31,25,90,.13), inset 0 1px 0 rgba(255,255,255,.82); }
.stTabs [data-baseweb="tab"] { min-height: 42px; padding: 10px 16px; font-size: 13px; transition: background .18s ease, color .18s ease, transform .18s ease; }
.stTabs [data-baseweb="tab"]:hover { color: var(--accent1); background: rgba(79,70,229,.08); }
.stTabs [aria-selected="true"] { box-shadow: 0 8px 18px rgba(79,70,229,.24) !important; }
.stTabs [data-baseweb="tab-highlight"] { display: none; }
.stButton > button, .stDownloadButton > button { min-height: 48px; border-radius: 14px; letter-spacing: .05px; }
.stButton > button:focus, .stDownloadButton > button:focus { outline: 3px solid rgba(14,165,233,.25); outline-offset: 2px; }
[data-testid="stFileUploader"] { padding: 10px; border-radius: 20px; box-shadow: 0 8px 24px rgba(31,25,90,.08); }
[data-testid="stFileUploader"] section { padding: 18px 12px; min-height: 126px; display: flex; align-items: center; justify-content: center; }
[data-testid="stFileUploader"] button { border-radius: 10px; }
[data-testid="stMetric"] { border-radius: 17px; }
.stAlert { border-radius: 15px; }
.stCaption, [data-testid="stCaptionContainer"] { color: var(--ink-faint); }
.stTextInput input, .stNumberInput input { min-height: 44px; }
[data-baseweb="select"] > div { min-height: 44px; border-radius: 12px; }
[data-testid="stExpander"] { border: 1px solid rgba(124,58,237,.14); border-radius: 16px; background: rgba(255,255,255,.38); overflow: hidden; }
[data-testid="stExpander"] summary { font-weight: 700; }
hr { border-color: rgba(79,70,229,.12); }
.sm-document-bar { display:flex; align-items:center; justify-content:space-between; gap:16px; flex-wrap:wrap; margin: 8px 0 17px; padding: 14px 18px; border-radius: 17px; background: rgba(255,255,255,.48); border: 1px solid rgba(255,255,255,.72); }
.sm-document-name { display:flex; align-items:center; gap:10px; min-width:0; color:var(--ink); font-size:20px; font-weight:850; letter-spacing:-.5px; }
.sm-document-name span:first-child { overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }
.sm-document-meta { color:var(--ink-faint); font-size:12px; font-weight:750; }
.sm-progress-dashboard { display:flex; align-items:center; justify-content:space-between; gap:26px; flex-wrap:wrap; padding:22px 26px; margin-top:16px; border-radius:26px; background:linear-gradient(115deg,rgba(255,255,255,.66),rgba(224,231,255,.48) 48%,rgba(253,242,255,.58)); border:1px solid rgba(255,255,255,.74); box-shadow:0 16px 42px rgba(31,25,90,.14), inset 0 1px 0 rgba(255,255,255,.86); }
.sm-progress-copy { flex:1; min-width:220px; }
.sm-progress-kicker { color:var(--accent1); font-size:11px; font-weight:800; letter-spacing:.7px; text-transform:uppercase; }
.sm-progress-title { color:var(--ink); font-size:22px; font-weight:850; margin-top:4px; }
.sm-progress-sub { color:var(--ink-faint); font-size:12px; margin-top:4px; }
.sm-progress-orb { --progress:0; position:relative; width:128px; height:128px; display:grid; place-items:center; border-radius:50%; background:conic-gradient(var(--accent1) 0 calc(var(--progress) * .7%), var(--accent2) calc(var(--progress) * .7%) calc(var(--progress) * 1%), rgba(148,163,184,.18) calc(var(--progress) * 1%) 100%); box-shadow:0 12px 28px rgba(79,70,229,.22); }
.sm-progress-orb::before { content:""; position:absolute; inset:10px; border-radius:50%; background:linear-gradient(145deg,rgba(255,255,255,.88),rgba(239,242,255,.72)); box-shadow:inset 0 2px 8px rgba(31,25,90,.08); }
.sm-progress-orb-content { position:relative; text-align:center; color:var(--ink); }
.sm-progress-number { display:block; font-size:27px; font-weight:850; letter-spacing:-1px; }
.sm-progress-label { display:block; color:var(--ink-faint); font-size:10px; font-weight:750; }
.sm-dashboard { margin:0 0 24px; padding:28px; border-radius:30px; background:linear-gradient(125deg,rgba(255,255,255,.72),rgba(224,231,255,.52) 46%,rgba(253,242,255,.62)); border:1px solid rgba(255,255,255,.8); box-shadow:0 20px 52px rgba(31,25,90,.16), inset 0 1px 0 rgba(255,255,255,.9); overflow:hidden; }
.sm-dashboard-heading { display:flex; justify-content:space-between; align-items:flex-start; gap:20px; flex-wrap:wrap; margin-bottom:24px; }
.sm-dashboard-heading h2 { margin:4px 0 4px; color:var(--ink); font-size:26px; letter-spacing:-.8px; }
.sm-dashboard-heading p { margin:0; color:var(--ink-faint); font-size:13px; }
.sm-dashboard-recommend { min-width:220px; padding:14px 16px; border-radius:17px; background:linear-gradient(135deg,rgba(79,70,229,.11),rgba(219,39,119,.10)); border:1px solid rgba(124,58,237,.16); }
.sm-dashboard-recommend span { display:block; color:var(--ink-faint); font-size:10px; font-weight:800; text-transform:uppercase; letter-spacing:.6px; margin-bottom:5px; }
.sm-dashboard-recommend b { color:var(--accent1); font-size:13px; }
.sm-dashboard-grid { display:grid; grid-template-columns:2fr repeat(4,1fr); gap:12px; }
.sm-dashboard-gauge { display:flex; align-items:center; gap:18px; padding:17px; min-height:128px; border-radius:21px; background:rgba(255,255,255,.5); border:1px solid rgba(255,255,255,.7); }
.sm-dashboard-gauge .sm-progress-orb { flex:0 0 auto; width:92px; height:92px; }
.sm-dashboard-gauge .sm-progress-orb::before { inset:8px; }
.sm-dashboard-gauge .sm-progress-number { font-size:21px; }
.sm-dashboard-card-kicker, .sm-dashboard-section-title { color:var(--accent1); font-size:10px; font-weight:850; letter-spacing:.7px; text-transform:uppercase; }
.sm-dashboard-gauge b { display:block; color:var(--ink); margin:4px 0; }
.sm-dashboard-gauge p { margin:0; color:var(--ink-faint); font-size:11px; line-height:1.4; }
.sm-dashboard-stat { display:flex; flex-direction:column; justify-content:center; min-height:128px; padding:15px; border-radius:21px; background:rgba(255,255,255,.46); border:1px solid rgba(255,255,255,.7); }
.sm-dashboard-stat span { font-size:22px; margin-bottom:8px; }
.sm-dashboard-stat b { color:var(--ink); font-size:23px; letter-spacing:-.7px; }
.sm-dashboard-stat small { color:var(--ink-faint); font-size:11px; margin-top:3px; }
.sm-dashboard-columns { display:grid; grid-template-columns:1fr 1fr; gap:14px; margin-top:18px; }
.sm-dashboard-section-title { margin:0 0 8px 3px; }
.sm-dashboard-row, .sm-dashboard-summary { display:flex; justify-content:space-between; gap:12px; padding:10px 12px; margin-top:6px; border-radius:12px; background:rgba(255,255,255,.42); color:var(--ink-soft); font-size:12px; }
.sm-dashboard-row span:last-child { color:var(--ink-faint); font-size:10px; }
.sm-dashboard-summary { display:block; }
.sm-dashboard-summary b { display:block; color:var(--ink); margin-bottom:3px; }
.sm-dashboard-summary span { display:block; color:var(--ink-faint); overflow-wrap:anywhere; word-break:break-word; white-space:normal; line-height:1.45; max-height:3.1em; overflow:hidden; }
.sm-dashboard-row span:first-child { min-width:0; overflow-wrap:anywhere; word-break:break-word; }
.sm-dashboard-empty { padding:12px; border-radius:12px; background:rgba(255,255,255,.32); color:var(--ink-faint); font-size:12px; }
.sm-dashboard-insights { display:flex; align-items:center; gap:8px; flex-wrap:wrap; margin-top:14px; padding:12px 14px; border-radius:16px; background:rgba(31,41,75,.045); border:1px solid rgba(124,58,237,.10); }
.sm-dashboard-insights span { padding:5px 10px; color:var(--ink-faint); font-size:11px; border-right:1px solid rgba(124,58,237,.12); }
.sm-dashboard-insights span:last-child { border-right:0; }
.sm-dashboard-insights b { color:var(--ink); margin-left:4px; }
.sm-dashboard-insights b.positive { color:#059669; }
.sm-dashboard-insights b.negative { color:#dc2626; }
.sm-dashboard-columns > div { min-width:0; }
.stApp .main .block-container { width:100%; max-width:1440px; }
.stTabs [data-baseweb="tab-list"] { overflow-x:auto; scrollbar-width:thin; }
.stTabs [data-baseweb="tab"] { white-space:nowrap; }
.stButton > button, .stDownloadButton > button { box-shadow:0 5px 14px rgba(31,25,90,.07); transition:transform .16s ease, box-shadow .16s ease; }
.stButton > button:hover, .stDownloadButton > button:hover { transform:translateY(-1px); box-shadow:0 8px 18px rgba(79,70,229,.16); }
.sm-sr-only { position:absolute !important; width:1px !important; height:1px !important; padding:0 !important; margin:-1px !important; overflow:hidden !important; clip:rect(0,0,0,0) !important; white-space:nowrap !important; border:0 !important; }
button:focus-visible, input:focus-visible, select:focus-visible, textarea:focus-visible, [tabindex="0"]:focus-visible { outline:3px solid #0ea5e9 !important; outline-offset:3px !important; box-shadow:0 0 0 5px rgba(14,165,233,.18) !important; }
label.sm-flip-inner { cursor:pointer; }
label.sm-flip-inner:focus-within { outline:3px solid #0ea5e9; outline-offset:5px; border-radius:20px; }
@media (prefers-reduced-motion: reduce) { *, *::before, *::after { animation-duration:.001ms !important; animation-iteration-count:1 !important; transition-duration:.001ms !important; scroll-behavior:auto !important; } }
.stMarkdown, [data-testid="stMarkdownContainer"] { overflow-wrap:anywhere; }
.stDataFrame, [data-testid="stMetric"] { border:1px solid rgba(124,58,237,.10); box-shadow:0 8px 22px rgba(31,25,90,.06); }
@media (max-width: 1050px) { .sm-dashboard-grid { grid-template-columns:repeat(2,1fr); } .sm-dashboard-gauge { grid-column:span 2; } }
@media (max-width: 640px) { .sm-dashboard { padding:20px 16px; } .sm-dashboard-grid, .sm-dashboard-columns { grid-template-columns:1fr; } .sm-dashboard-gauge { grid-column:auto; } .sm-dashboard-insights { display:grid; grid-template-columns:1fr 1fr; } .sm-dashboard-insights span { border-right:0; border-bottom:1px solid rgba(124,58,237,.10); } .sm-dashboard-insights span:last-child { grid-column:span 2; } }
.sm-footer { opacity:.8; }
.sm-audio-panel { margin-top:20px; padding:20px; border-radius:20px; background:linear-gradient(135deg,rgba(224,231,255,.72),rgba(253,242,255,.72)); border:1px solid rgba(124,58,237,.16); box-shadow:0 12px 28px rgba(31,25,90,.1), inset 0 1px 0 rgba(255,255,255,.8); }
.sm-audio-title { color:var(--ink); font-size:16px; font-weight:750; margin-bottom:5px; }
.sm-audio-copy { color:var(--ink-faint); font-size:12px; margin-bottom:12px; }
@media (max-width: 900px) { .main .block-container { padding: 1.25rem 1rem 3rem; } .stTabs [data-baseweb="tab-list"] { position: static; } }
@media (max-width: 640px) { .sm-hero { padding: 40px 18px 34px; border-radius: 24px; } .sm-hero h1 { font-size: 38px; } .sm-hero p.sub { font-size: 16px; } .sm-feature { min-height: 0; } .sm-document-name { font-size: 17px; } }
</style> 
""") 

html("""
<style>
/* PERFORMANCE PASS: preserve the glass look without continuous GPU-heavy motion */
.sm-hero, .sm-card, .sm-feature, .sm-metric, .sm-document-bar { animation:none !important; }
.sm-hero::after { animation:none !important; filter:none !important; opacity:.32; }
.sm-metric, .stButton > button, .stDownloadButton > button, .stTabs [data-baseweb="tab"] { transition:none !important; }
.sm-metric:hover, .stButton > button:hover, .stDownloadButton > button:hover { transform:none !important; }
.stApp > header, [data-testid="stHeader"] { backdrop-filter:blur(14px) saturate(135%) !important; -webkit-backdrop-filter:blur(14px) saturate(135%) !important; }
[data-testid="stToolbar"], [data-testid="stStatusWidget"] { backdrop-filter:blur(10px) !important; -webkit-backdrop-filter:blur(10px) !important; }
.sm-card, .sm-dashboard, .sm-feature, .sm-metric, .sm-document-bar, .sm-progress-dashboard { will-change:auto !important; }
@media (max-width: 700px) {
  .main .block-container { padding: .85rem .7rem 5.5rem !important; }
  .sm-hero { margin: 0 -2px 16px; padding: 32px 18px 28px !important; }
  .sm-dashboard { padding: 18px 14px !important; border-radius: 22px; }
  .sm-dashboard-grid, .sm-dashboard-columns { grid-template-columns: 1fr !important; }
  .sm-dashboard-gauge { grid-column:auto !important; }
  .sm-document-bar { padding: 11px 12px; }
  .sm-document-name { max-width: 100%; font-size: 16px; }
  .sm-flip-grid { display:flex !important; grid-template-columns:none !important; gap:14px !important; overflow-x:auto; overscroll-behavior-x:contain; scroll-snap-type:x mandatory; padding:2px 4px 14px; }
  .sm-flip-grid > .sm-flip-card { flex:0 0 86vw; scroll-snap-align:center; }
  .sm-flip-card, .sm-flip-inner { min-height: 270px !important; }
  .stButton > button, .stDownloadButton > button { min-height: 52px !important; font-size: 15px !important; }
  [data-testid="stAudio"] audio { width: 100% !important; min-height: 48px; }
  .stTabs [data-baseweb="tab-list"] { position: fixed !important; left: 0; right: 0; bottom: 0; top: auto !important; z-index: 1000; padding: 6px 4px !important; border-radius: 18px 18px 0 0 !important; overflow-x: auto; background: rgba(255,255,255,.92) !important; backdrop-filter: blur(18px); box-shadow: 0 -8px 24px rgba(31,25,90,.16) !important; }
  .stTabs [data-baseweb="tab"] { min-width: 78px; min-height: 46px; padding: 8px 10px !important; font-size: 11px !important; }
}
.sm-privacy-note { margin-top: 8px; padding: 10px 12px; border-radius: 12px; color: var(--ink-faint); background: rgba(79,70,229,.06); border: 1px solid rgba(79,70,229,.12); font-size: 11px; line-height: 1.45; }
</style>
""")
 
# ============================================================ 
# GROQ CLIENT 
# ============================================================ 
 
MODEL_CANDIDATES = ["openai/gpt-oss-120b"] 
 
def get_groq_client(): 
    api_key = os.environ.get("GROQ_API_KEY") 
    if not api_key: 
        try: 
            api_key = st.secrets.get("GROQ_API_KEY") 
        except Exception: 
            api_key = None 
    if not api_key: 
        st.error("Groq API key is not configured. Set GROQ_API_KEY as an env var or in Streamlit Secrets.") 
        return None 
    return Groq(api_key=api_key) 
 
 
@st.cache_data(ttl=3600, max_entries=128, show_spinner=False)
def cached_groq_response(prompt, system_message, max_tokens):
    client = get_groq_client()
    if client is None:
        return ""
    for model in MODEL_CANDIDATES:
        try:
            messages = []
            if system_message:
                messages.append({"role": "system", "content": system_message})
            messages.append({"role": "user", "content": prompt})
            response = client.chat.completions.create(model=model, messages=messages, temperature=0.2, max_tokens=max_tokens)
            return response.choices[0].message.content
        except Exception:
            continue
    return ""


def ask_groq(prompt, system_message=None, max_tokens=2000): 
    """Cached AI request wrapper to avoid duplicate generation on Streamlit reruns."""
    cached = cached_groq_response(prompt, system_message or "", max_tokens)
    if cached:
        return cached
    client = get_groq_client() 
    if client is None: 
        return "" 
    messages = [] 
    if system_message: 
        messages.append({"role": "system", "content": system_message}) 
    messages.append({"role": "user", "content": prompt}) 
    last_error = None 
    for model in MODEL_CANDIDATES: 
        try: 
            response = client.chat.completions.create( 
                model=model, messages=messages, temperature=0.2, max_tokens=max_tokens 
            ) 
            return response.choices[0].message.content 
        except Exception as e: 
            last_error = e 
            continue 
    st.error(f"Groq API error: {last_error}") 
    return "" 


def generate_speech_audio(text, language="en"):
    """Create MP3 audio bytes for a study passage using the optional gTTS backend."""
    if gTTS is None:
        st.error("Text-to-speech is not installed. Add gTTS to requirements.txt and redeploy.")
        return None
    try:
        audio_buffer = BytesIO()
        gTTS(text=text, lang=language, slow=False).write_to_fp(audio_buffer)
        return audio_buffer.getvalue()
    except Exception as e:
        st.error(f"Could not create audio: {e}")
        return None


def clean_speech_text(text):
    """Remove markdown, document markers, and formatting artifacts before TTS."""
    text = re.sub(r"```.*?```", " ", text, flags=re.DOTALL)
    text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s*[-*+]\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s*\d+[.)]\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"\[/?(?:PAGE|SLIDE|TABLE)[^\]]*\]", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"https?://\S+", " ", text)
    text = re.sub(r"[*_`~]", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def prepare_speech_text(text, language_name):
    """Clean English summaries and translate them for non-English voices."""
    cleaned = clean_speech_text(text)
    if language_name == "English":
        return cleaned
    prompt = f"Translate the study summary below into natural, clear {language_name} for spoken audio. Preserve all factual meaning. Return only the translated study content without an introduction, headings, markdown, bullets, symbols, or commentary.\n\nSUMMARY:\n{cleaned[:12000]}"
    translated = ask_groq(prompt, max_tokens=3500)
    return clean_speech_text(translated) if translated else cleaned
 
 
# ============================================================
# DOCUMENT EXTRACTION / CLEANING / CHUNKING
# ============================================================

SUPPORTED_TYPES = ["pdf", "docx", "pptx", "txt", "md"]
SUPPORTED_LABELS = {"pdf": "PDF", "docx": "Word", "pptx": "PowerPoint", "txt": "Text", "md": "Markdown", "workspace": "Workspace"}


def extract_pdf_text(uploaded_file):
    try:
        uploaded_file.seek(0)
        reader = PdfReader(uploaded_file)
        pages = []
        for n, page in enumerate(reader.pages, start=1):
            text = page.extract_text()
            if text:
                pages.append(f"\n[PAGE {n}]\n{text}")
        extracted = "".join(pages)
        if extracted.strip():
            return extracted

        if pytesseract is None or convert_from_bytes is None:
            st.warning("This PDF appears to be scanned, but OCR support is not installed. Add pytesseract, pdf2image, and Poppler to enable scanned-document reading.")
            return ""

        st.info("No selectable text was found. Studient AI is using OCR to read the scanned pages...")
        uploaded_file.seek(0)
        images = convert_from_bytes(uploaded_file.getvalue(), dpi=220, fmt="png", thread_count=2)
        ocr_pages = []
        for number, image in enumerate(images, start=1):
            ocr_text = pytesseract.image_to_string(image, config="--psm 6")
            if ocr_text.strip():
                ocr_pages.append(f"\n[OCR PAGE {number}]\n{ocr_text}")
        return "".join(ocr_pages)
    except Exception as e:
        st.error(f"Could not read the PDF or run OCR: {e}")
        return ""


def extract_docx_text(uploaded_file):
    if DocxDocument is None:
        st.error("Word support is not installed. Run: pip install python-docx")
        return ""
    try:
        uploaded_file.seek(0)
        document = DocxDocument(uploaded_file)
        parts = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)
        for table_index, table in enumerate(document.tables, start=1):
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            if rows:
                parts.append(f"\n[TABLE {table_index}]\n" + "\n".join(rows))
        return "\n".join(parts)
    except Exception as e:
        st.error(f"Could not read the Word document: {e}")
        return ""


def extract_pptx_text(uploaded_file):
    if Presentation is None:
        st.error("PowerPoint support is not installed. Run: pip install python-pptx")
        return ""
    try:
        uploaded_file.seek(0)
        presentation = Presentation(uploaded_file)
        slides = []
        for number, slide in enumerate(presentation.slides, start=1):
            slide_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_parts.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        slide_parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            if slide_parts:
                slides.append(f"\n[SLIDE {number}]\n" + "\n".join(slide_parts))
        return "\n".join(slides)
    except Exception as e:
        st.error(f"Could not read the PowerPoint presentation: {e}")
        return ""


def extract_plain_text(uploaded_file):
    try:
        uploaded_file.seek(0)
        return uploaded_file.read().decode("utf-8", errors="replace")
    except Exception as e:
        st.error(f"Could not read the text file: {e}")
        return ""


def extract_document_text(uploaded_file):
    extension = uploaded_file.name.rsplit(".", 1)[-1].lower() if "." in uploaded_file.name else ""
    extractors = {"pdf": extract_pdf_text, "docx": extract_docx_text, "pptx": extract_pptx_text, "txt": extract_plain_text, "md": extract_plain_text}
    extractor = extractors.get(extension)
    if extractor is None:
        st.error("Unsupported document format. Please upload a PDF, Word document, or PowerPoint presentation.")
        return "", extension
    return extractor(uploaded_file), extension
def clean_text(text): 
    text = text.replace("\x00", " ") 
    text = re.sub(r"[ \t]+", " ", text) 
    text = re.sub(r"\n{3,}", "\n\n", text) 
    return text.strip() 
 
 
def chunk_text(text, chunk_size=7000): 
    words = text.split() 
    chunks, current, length = [], [], 0 
    for w in words: 
        current.append(w) 
        length += len(w) + 1 
        if length >= chunk_size: 
            chunks.append(" ".join(current)) 
            current, length = [], 0 
    if current: 
        chunks.append(" ".join(current)) 
    return chunks 
 
 
@st.cache_data(max_entries=128, show_spinner=False)
def get_relevant_chunks(text, query, max_chunks=8, max_context_chars=12000):
    """
    Select the most relevant chunks while enforcing a hard context-size limit.
    This prevents Groq 413 / TPM errors on large PDFs and full books.
    """

    chunks = chunk_text(text, chunk_size=7000)

    if not chunks:
        return ""

    query_words = set(re.findall(r"\b[a-zA-Z0-9]{3,}\b", query.lower()))

    scored = []

    for i, chunk in enumerate(chunks):
        chunk_words = set(
            re.findall(r"\b[a-zA-Z0-9]{3,}\b", chunk.lower())
        )

        overlap = len(query_words & chunk_words)

        # Small bonus for chunks containing important academic terms
        important_terms = {
            "definition",
            "concept",
            "important",
            "process",
            "theory",
            "principle",
            "example",
            "formula",
            "function",
            "cause",
            "effect",
            "classification",
        }

        bonus = len(chunk_words & important_terms)

        score = overlap + (bonus * 0.5)

        scored.append((score, i, chunk))

    # Highest relevance first
    scored.sort(key=lambda x: x[0], reverse=True)

    selected = []
    total_chars = 0

    for score, index, chunk in scored[:max_chunks]:

        remaining = max_context_chars - total_chars

        if remaining <= 500:
            break

        if len(chunk) <= remaining:
            selected.append((index, chunk))
            total_chars += len(chunk)
        else:
            # Take only the portion that fits
            selected.append((index, chunk[:remaining]))
            total_chars += remaining
            break

    # Restore original PDF order
    selected.sort(key=lambda x: x[0])

    return "\n\n".join(chunk for _, chunk in selected)
 
 
# ============================================================ 
# JSON QUIZ GENERATION (structured, not regex-parsed) 
# ============================================================ 
 
QUIZ_SYSTEM_MSG = ( 
    "You are a quiz generation engine. Respond with ONLY valid JSON — " 
    "no markdown, no code fences, no commentary. The JSON must be a list of objects " 
    'with keys: "question" (string), "options" (object with keys A,B,C,D), ' 
    '"correct" (one of "A","B","C","D"), "explanation" (string), ' 
    '"topic" (short 2-4 word label). Base everything strictly on the material given.' 
) 
 
def _strip_json_fences(raw): 
    raw = raw.strip() 
    raw = re.sub(r"^```(json)?", "", raw).strip() 
    raw = re.sub(r"```$", "", raw).strip() 
    return raw 
 
 
def generate_quiz_questions(text, count, difficulty, category="Mixed"):
    category_focus = {"Mixed": "important exam concepts facts definitions processes", "Definitions": "definitions terms meanings", "Processes": "processes mechanisms steps cause effect", "Comparisons": "comparisons differences similarities", "Facts": "important facts values examples", "Exam-focused": "high priority exam concepts likely questions"}.get(category, "important concepts")
    context = get_relevant_chunks(text, category_focus, max_chunks=10) 
    prompt = f"""Create exactly {count} multiple-choice questions at {difficulty} difficulty in the {category} category
from the study material below. Return ONLY the JSON array described in the system message. 
 
STUDY MATERIAL: 
{context} 
""" 
    raw = ask_groq(prompt, system_message=QUIZ_SYSTEM_MSG, max_tokens=4000) 
    if not raw: 
        return None 
    try: 
        data = json.loads(_strip_json_fences(raw)) 
        cleaned = [q for q in data if all(k in q for k in ("question", "options", "correct", "explanation", "topic"))
                   and all(k in q["options"] for k in ("A", "B", "C", "D"))]
        for q in cleaned:
            correct_text = q["options"][q["correct"]]
            options = list(q["options"].values())
            random.shuffle(options)
            q["options"] = dict(zip(("A", "B", "C", "D"), options))
            q["correct"] = next(letter for letter, value in q["options"].items() if value == correct_text)
        return cleaned or None 
    except Exception as e: 
        st.error(f"Couldn't parse the generated quiz as JSON: {e}") 
        with st.expander("Raw model output (for debugging)"): 
            st.code(raw) 
        return None 
 
 
# ============================================================ 
# FLASHCARD GENERATION -- separate from MCQs on purpose. 
# MCQ "correct answer" text is deliberately short (it's one of 
# four options); a flashcard back needs a real explanation, so 
# this uses its own prompt/schema instead of reusing MCQ output. 
# ============================================================ 
 
# ============================================================
# FLASHCARD GENERATION -- visual artifacts + flip cards
# ============================================================

FLASHCARD_SYSTEM_MSG = (
    "You are an educational visual flashcard generation engine. "
    "Respond with ONLY valid JSON — no markdown, no code fences, no commentary. "
    "The JSON must be a list of objects with these keys: "
    '"front" (a concise question or concept, ideally under 12 words), '
    '"topic" (a short 2-4 word label), '
    '"visual_type" (one of: definition, process, comparison, timeline, formula, fact), '
    '"visual_title" (a short title for the visual), '
    '"visual_items" (an array of 2-6 short strings containing the key information needed '
    'to understand the concept visually). '
    "Choose the visual_type that best represents the concept. "
    "For process use ordered steps. "
    "For comparison use two sides separated with 'VS'. "
    "For timeline use chronological events. "
    "For formula use the important equation or relationship. "
    "For definition use the central concept plus supporting points. "
    "For fact use the most important fact or relationship. "
    "Keep visual_items concise and easy to display in a diagram. "
    "Base everything strictly on the study material. Do not invent facts."
)


def generate_flashcards(text, count, difficulty):
    context = get_relevant_chunks(
        text,
        "definitions concepts important terms facts explanations processes",
        max_chunks=8,
        max_context_chars=12000
    )

    prompt = f"""Create exactly {count} visual study flashcards at {difficulty} depth.

Each flashcard must turn an important concept from the material into a visual learning artifact.

Requirements:
- "front" = concise question or concept, ideally under 12 words.
- "topic" = short 2-4 word category.
- "visual_type" = definition, process, comparison, timeline, formula, or fact.
- "visual_title" = short title.
- "visual_items" = 2-6 concise pieces of information.
- Do NOT create generic or unrelated visuals.
- Use a process visual when the material describes steps.
- Use comparison when two concepts are contrasted.
- Use timeline when events/stages have chronological order.
- Use formula when an equation or relationship is important.
- Use definition for concepts that are best understood through a central idea.
- Use fact for an important relationship, value, or memorable fact.

Return ONLY the JSON array described in the system message.

STUDY MATERIAL:
{context}
"""

    raw = ask_groq(
        prompt,
        system_message=FLASHCARD_SYSTEM_MSG,
        max_tokens=2500
    )

    if not raw:
        return None

    try:
        data = json.loads(_strip_json_fences(raw))

        cleaned = []

        for card in data:
            if not isinstance(card, dict):
                continue

            required = [
                "front",
                "topic",
                "visual_type",
                "visual_title",
                "visual_items"
            ]

            if not all(k in card for k in required):
                continue

            if card["visual_type"] not in {
                "definition",
                "process",
                "comparison",
                "timeline",
                "formula",
                "fact"
            }:
                card["visual_type"] = "definition"

            if not isinstance(card["visual_items"], list):
                continue

            items = [
                str(item).strip()
                for item in card["visual_items"]
                if str(item).strip()
            ]

            if not items:
                continue

            card["visual_items"] = items[:6]

            cleaned.append(card)

        return cleaned or None

    except Exception as e:
        st.error(f"Couldn't parse the generated flashcards as JSON: {e}")

        with st.expander("Raw model output (for debugging)"):
            st.code(raw)

        return None


def render_visual_artifact(card):
    esc = _html_escape_lib.escape

    visual_type = card["visual_type"]
    title = esc(card["visual_title"])
    items = [esc(str(x)) for x in card["visual_items"]]

    pieces = [
        '<div class="sm-visual-artifact">',
        f'<div class="sm-visual-title">{title}</div>'
    ]

    # ---------------- DEFINITION ----------------
    if visual_type == "definition":
        core = items[0]

        pieces.append(
            '<div class="sm-visual-definition">'
            f'<div class="sm-visual-core">{core}</div>'
            '</div>'
        )

        if len(items) > 1:
            pieces.append('<div class="sm-visual-items">')

            for item in items[1:]:
                pieces.append(
                    f'<div class="sm-visual-item">• {item}</div>'
                )

            pieces.append('</div>')

    # ---------------- PROCESS ----------------
    elif visual_type == "process":
        pieces.append('<div class="sm-process">')

        for index, item in enumerate(items):
            pieces.append(
                f'<div class="sm-process-step">'
                f'{index + 1}. {item}'
                f'</div>'
            )

            if index < len(items) - 1:
                pieces.append('<div class="sm-arrow">↓</div>')

        pieces.append('</div>')

    # ---------------- COMPARISON ----------------
    elif visual_type == "comparison":
        left_items = []
        right_items = []

        for index, item in enumerate(items):
            if index % 2 == 0:
                left_items.append(item)
            else:
                right_items.append(item)

        pieces.append('<div class="sm-comparison">')

        pieces.append(
            '<div class="sm-comparison-column">'
            '<div class="sm-comparison-title">Concept A</div>'
        )

        for item in left_items:
            pieces.append(
                f'<div class="sm-comparison-item">• {item}</div>'
            )

        pieces.append('</div>')

        pieces.append(
            '<div class="sm-comparison-column">'
            '<div class="sm-comparison-title">Concept B</div>'
        )

        for item in right_items:
            pieces.append(
                f'<div class="sm-comparison-item">• {item}</div>'
            )

        pieces.append('</div>')
        pieces.append('</div>')

    # ---------------- TIMELINE ----------------
    elif visual_type == "timeline":
        pieces.append('<div class="sm-timeline">')

        for index, item in enumerate(items):
            pieces.append(
                '<div class="sm-timeline-item">'
                f'<div class="sm-timeline-number">{index + 1}</div>'
                f'<div class="sm-timeline-text">{item}</div>'
                '</div>'
            )

        pieces.append('</div>')

    # ---------------- FORMULA ----------------
    elif visual_type == "formula":
        formula = "<br>".join(items)

        pieces.append(
            f'<div class="sm-formula">{formula}</div>'
        )

    # ---------------- FACT ----------------
    elif visual_type == "fact":
        fact = " • ".join(items)

        pieces.append(
            f'<div class="sm-fact">{fact}</div>'
        )

    pieces.append('</div>')

    return "".join(pieces)


def render_flashcards(cards):
    """Render interactive click-to-flip visual flashcards."""

    esc = _html_escape_lib.escape

    pieces = ['<div class="sm-flip-grid">']

    for i, card in enumerate(cards):

        front = esc(str(card["front"]))
        topic = esc(str(card["topic"]))

        visual = render_visual_artifact(card)

        pieces.append(f"""
        <div class="sm-flip-card">

          <input
            type="checkbox"
            id="flipcard_{i}"
          >

          <label
            for="flipcard_{i}"
            class="sm-flip-inner"
            tabindex="0"
            aria-label="Flashcard {i + 1}: {front}. Press space or enter to reveal the answer."
          >

            <!-- FRONT -->
            <div class="sm-flip-front">

              <span class="sm-sr-only">Flashcard {i + 1}. Question: {front}. Press the spacebar or Enter to flip this card.</span>

              <div>
                <div class="sm-quiz-topic">
                  {topic}
                </div>

                <div class="sm-flip-q">
                  {front}
                </div>
              </div>

              <div class="sm-flip-hint">
                👆 Click to reveal visual answer
              </div>

            </div>

            <!-- BACK -->
            <div class="sm-flip-back">

              <div class="sm-flip-a-label">
                Visual Explanation
              </div>

              {visual}

              <div class="sm-flip-back-note">
                Based strictly on your uploaded study material.
              </div>

              <div class="sm-flip-hint">
                👆 Click to flip back
              </div>

            </div>

          </label>

        </div>
        """)

    pieces.append("</div>")

    html("".join(pieces))
 
 

# ============================================================
# INTERACTIVE MIND MAP GENERATION
# ============================================================

MIND_MAP_SYSTEM_MSG = (
    "You are an educational mind map architect. Respond with ONLY valid JSON — no markdown, "
    "no code fences, no commentary. Return one object with keys: title (string) and children (array). "
    "Every node is an object with label (short string) and children (array of nodes). "
    "Create a useful hierarchy from the supplied study material: 4-7 major branches, each with 2-5 "
    "supporting concepts, and deeper nodes only when genuinely useful. Keep labels concise. "
    "Base every node strictly on the material and do not invent facts."
)


def _clean_mind_map_node(node, depth=0):
    if not isinstance(node, dict):
        return None
    label = str(node.get("label", "")).strip()
    if not label:
        return None
    raw_children = node.get("children", [])
    children = []
    if depth < 4 and isinstance(raw_children, list):
        for child in raw_children[:8]:
            cleaned = _clean_mind_map_node(child, depth + 1)
            if cleaned:
                children.append(cleaned)
    return {"label": label[:120], "children": children}


def generate_mind_map(text):
    context = get_relevant_chunks(
        text,
        "main topics concepts definitions processes examples relationships causes effects overview",
        max_chunks=10,
        max_context_chars=14000,
    )
    prompt = f"""Build a hierarchical interactive mind map for the study material below.

The root title should describe the overall document or subject. Organize the most important ideas into
clear branches suitable for revision. Prefer meaningful relationships over a long list of isolated facts.
Return ONLY the JSON object described in the system message.

STUDY MATERIAL:
{context}
"""
    raw = ask_groq(prompt, system_message=MIND_MAP_SYSTEM_MSG, max_tokens=3500)
    if not raw:
        return None
    try:
        data = json.loads(_strip_json_fences(raw))
        root_label = str(data.get("title", "Study Material")).strip() if isinstance(data, dict) else "Study Material"
        raw_children = data.get("children", []) if isinstance(data, dict) else []
        root = {"label": root_label[:120] or "Study Material", "children": []}
        if isinstance(raw_children, list):
            for child in raw_children[:8]:
                cleaned = _clean_mind_map_node(child, 1)
                if cleaned:
                    root["children"].append(cleaned)
        return root if root["children"] else None
    except Exception as e:
        st.error(f"Couldn't parse the generated mind map as JSON: {e}")
        with st.expander("Raw model output (for debugging)"):
            st.code(raw)
        return None


def _render_mind_map_nodes(nodes, root=False):
    esc = _html_escape_lib.escape
    parts = []
    for node in nodes:
        label = esc(str(node.get("label", "")))
        children = node.get("children") or []
        node_class = "sm-mindmap-root" if root else ""
        parts.append(f'<li class="{node_class}">')
        if children:
            parts.append('<details open class="sm-mindmap-details">')
            parts.append(f'<summary class="sm-mindmap-node" title="Click to collapse or expand this branch"><span class="sm-mindmap-toggle">−</span><span class="sm-mindmap-dot"></span><span>{label}</span></summary>')
            parts.append('<ul class="sm-mindmap-children">')
            parts.append(_render_mind_map_nodes(children))
            parts.append('</ul></details>')
        else:
            parts.append(f'<div class="sm-mindmap-node" title="Leaf concept"><span class="sm-mindmap-dot"></span><span>{label}</span></div>')
        parts.append('</li>')
    return "".join(parts)


def render_mind_map(mind_map):
    esc = _html_escape_lib.escape
    root = {"label": mind_map.get("label", "Study Material"), "children": mind_map.get("children", [])}
    parts = ['<div class="sm-mindmap-shell">']
    parts.append('<div class="sm-mindmap-toolbar"><div><b>Interactive concept map</b><div class="sm-mindmap-note">Follow a branch, then collapse it to focus your revision.</div></div><div class="sm-mindmap-note">Click any branch label to focus • scroll horizontally on smaller screens</div></div>')
    parts.append('<ul class="sm-mindmap-tree">')
    parts.append(_render_mind_map_nodes([root], root=True))
    parts.append('</ul></div>')
    html("".join(parts))

# ============================================================ 
# SCORE HISTORY (local JSON — resets on Streamlit Cloud restart) 
# ============================================================ 
 
HISTORY_FILE = "study_history.json" 
SUMMARY_HISTORY_FILE = "summary_history.json"
REVIEW_STATE_FILE = "flashcard_review_state.json"


def privacy_secret():
    """Read a deployment secret without ever displaying it in the UI."""
    return os.environ.get("STUDIENT_ENCRYPTION_KEY") or os.environ.get("APP_SECRET") or ""


def encryption_cipher():
    secret = privacy_secret()
    if Fernet is None or not secret:
        return None
    key = base64.urlsafe_b64encode(hashlib.sha256(secret.encode("utf-8")).digest())
    return Fernet(key)


def read_private_json(path, default):
    if not os.path.exists(path):
        return default
    try:
        raw = open(path, "rb").read()
        cipher = encryption_cipher()
        if cipher:
            try:
                raw = cipher.decrypt(raw)
            except InvalidToken:
                pass  # Support legacy unencrypted files during migration.
        return json.loads(raw.decode("utf-8"))
    except Exception:
        return default


def write_private_json(path, value):
    raw = json.dumps(value, indent=2, ensure_ascii=False).encode("utf-8")
    cipher = encryption_cipher()
    if cipher:
        raw = cipher.encrypt(raw)
    with open(path, "wb") as handle:
        handle.write(raw)


def secure_delete_file(path):
    if not os.path.exists(path):
        return
    try:
        size = os.path.getsize(path)
        with open(path, "r+b", buffering=0) as handle:
            handle.write(secrets.token_bytes(size))
            handle.flush()
            os.fsync(handle.fileno())
        os.remove(path)
    except OSError:
        try:
            os.remove(path)
        except OSError:
            pass


def delete_local_study_data():
    for path in [HISTORY_FILE, SUMMARY_HISTORY_FILE, REVIEW_STATE_FILE]:
        secure_delete_file(path)
    for key in ["pdf_text", "document_name", "workspace_documents", "summary_text", "flashcards", "mind_map", "study_plan", "generated_material", "tutor_messages"]:
        if key in st.session_state:
            st.session_state[key] = {} if key == "workspace_documents" else [] if key == "tutor_messages" else None if key not in ["pdf_text", "document_name"] else ""


def make_user_data_export():
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as pack:
        pack.writestr("workspace.txt", st.session_state.get("pdf_text", ""))
        pack.writestr("study-history.json", json.dumps(load_history(), indent=2))
        pack.writestr("summary-history.json", json.dumps(load_summary_history(), indent=2))
        pack.writestr("flashcard-review-state.json", json.dumps(st.session_state.get("review_state", {}), indent=2))
        if st.session_state.get("summary_text"):
            pack.writestr("current-summary.md", st.session_state.summary_text)
        if st.session_state.get("study_plan"):
            pack.writestr("study-plan.md", st.session_state.study_plan)
    return buffer.getvalue()


def apply_retention_policy(days):
    if not days:
        return
    cutoff = datetime.now() - timedelta(days=int(days))
    history = [item for item in load_history() if item.get("timestamp", "") >= cutoff.isoformat()]
    summaries = [item for item in load_summary_history() if item.get("timestamp", "") >= cutoff.isoformat()]
    write_private_json(HISTORY_FILE, history)
    write_private_json(SUMMARY_HISTORY_FILE, summaries)
 
def load_history(): 
    return read_private_json(HISTORY_FILE, [])
 
 
def save_attempt(document_name, score, total, topic_results): 
    if st.session_state.get("privacy_mode"):
        return []
    history = load_history() 
    history.append({ 
        "timestamp": datetime.now().isoformat(timespec="seconds"), 
        "document": document_name, "score": score, "total": total, "topics": topic_results, 
        "duration_minutes": round(max(0, (time.time() - (st.session_state.get("test_started_at") or time.time())) / 60), 1),
    }) 
    try: 
        write_private_json(HISTORY_FILE, history)
    except Exception as e: 
        st.warning(f"Couldn't save this attempt to history: {e}") 
    return history 


def make_pdf_bytes(title, sections):
    if SimpleDocTemplate is None:
        return None
    buffer = BytesIO()
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 0.18 * inch)]
    for heading, body in sections:
        story.append(Paragraph(str(heading), styles["Heading2"]))
        for paragraph in str(body).split("\n"):
            if paragraph.strip():
                story.append(Paragraph(_html_escape_lib.escape(paragraph.strip()), styles["BodyText"]))
                story.append(Spacer(1, 0.08 * inch))
    SimpleDocTemplate(buffer, pagesize=letter, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42).build(story)
    return buffer.getvalue()


def make_flashcards_pdf(cards):
    sections = []
    for index, card in enumerate(cards, 1):
        answer = "\n".join(card.get("visual_items", []))
        sections.append((f"Card {index}: {card.get('front', '')}", f"Topic: {card.get('topic', '')}\n{card.get('visual_title', '')}\n{answer}"))
    return make_pdf_bytes("Studient AI Flashcards", sections)


def make_questions_docx(questions, title="Studient AI Questions"):
    if Document is None:
        return None
    document = Document()
    document.add_heading(title, 0)
    for index, question in enumerate(questions, 1):
        document.add_heading(f"Question {index}", level=2)
        if isinstance(question, dict):
            document.add_paragraph(question.get("question", ""))
            for letter, option in question.get("options", {}).items():
                document.add_paragraph(f"{letter}) {option}")
            document.add_paragraph(f"Answer: {question.get('correct', '')}")
            document.add_paragraph(f"Explanation: {question.get('explanation', '')}")
        else:
            document.add_paragraph(str(question))
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_mindmap_png(mind_map):
    if Image is None:
        return None
    lines = []
    def walk(node, depth=0):
        lines.append("  " * depth + ("• " if depth else "") + str(node.get("label", "")))
        for child in node.get("children", []) or []:
            walk(child, depth + 1)
    walk(mind_map)
    width, line_height = 1400, 34
    height = max(220, 70 + len(lines) * line_height)
    image = Image.new("RGB", (width, height), (246, 247, 255))
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((24, 24, width - 24, height - 24), radius=24, fill=(255, 255, 255), outline=(124, 58, 237), width=3)
    y = 48
    for line in lines:
        draw.text((54, y), line[:115], fill=(31, 41, 75))
        y += line_height
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def make_study_pack():
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as pack:
        pack.writestr("workspace.txt", st.session_state.get("pdf_text", ""))
        if st.session_state.get("summary_text"):
            pack.writestr("summary.md", st.session_state.summary_text)
        if st.session_state.get("flashcards"):
            csv_buffer = StringIO()
            writer = csv.writer(csv_buffer)
            writer.writerow(["Front", "Topic", "Visual type", "Visual title", "Visual items"])
            for card in st.session_state.flashcards:
                writer.writerow([card.get("front", ""), card.get("topic", ""), card.get("visual_type", ""), card.get("visual_title", ""), " | ".join(card.get("visual_items", []))])
            pack.writestr("flashcards.csv", csv_buffer.getvalue())
        if st.session_state.get("mind_map"):
            pack.writestr("mind-map.json", json.dumps(st.session_state.mind_map, indent=2, ensure_ascii=False))
        pack.writestr("study-history.json", json.dumps(load_history(), indent=2))
    return buffer.getvalue()


def analytics_snapshot(document_name):
    history = [item for item in load_history() if item.get("document") == document_name]
    scores = [round(item["score"] / item["total"] * 100) for item in history if item.get("total")]
    topic_totals = defaultdict(lambda: {"correct": 0, "total": 0})
    for attempt in history:
        for topic, stats in attempt.get("topics", {}).items():
            topic_totals[topic]["correct"] += stats.get("correct", 0)
            topic_totals[topic]["total"] += stats.get("correct", 0) + stats.get("wrong", 0)
    topic_accuracy = {topic: round(data["correct"] / data["total"] * 100) if data["total"] else 0 for topic, data in topic_totals.items()}
    cards = st.session_state.get("flashcards") or []
    mastered = sum(1 for card in cards if get_review_card(card).get("box", 1) >= 5)
    dates = {item.get("timestamp", "")[:10] for item in history if item.get("timestamp")}
    avg = round(sum(scores) / len(scores)) if scores else 0
    improvement = scores[-1] - scores[0] if len(scores) > 1 else 0
    weakest = min(topic_accuracy, key=topic_accuracy.get) if topic_accuracy else "Complete a practice test"
    return {"attempts": len(history), "average": avg, "scores": scores, "improvement": improvement, "topic_accuracy": topic_accuracy, "mastered": mastered, "cards": len(cards), "questions": sum(item.get("total", 0) for item in history), "sessions": len(dates), "weakest": weakest}


def generate_study_plan(exam_date, target_grade, hours_per_day, weak_subjects, schedule):
    days_left = max(1, (exam_date - date.today()).days)
    context = get_relevant_chunks(st.session_state.pdf_text, weak_subjects or "all major topics", max_chunks=8)
    prompt = f"""Create a practical personalized study plan grounded in the workspace material below.
Exam date: {exam_date} ({days_left} days from today)
Target grade: {target_grade}
Available study time: {hours_per_day} hours per day
Weak subjects or topics: {weak_subjects or 'Identify from the material'}
Preferred schedule: {schedule}

Include: a clear phase overview, daily tasks, revision blocks, practice tests, flashcard reviews, rest/buffer days, and a final revision plan. Keep tasks realistic for the available time. Use markdown headings and a day-by-day table where useful.

WORKSPACE MATERIAL:
{context}"""
    return ask_groq(prompt, max_tokens=4200)


def tutor_reply(question, level, mode, conversation):
    context = get_relevant_chunks(st.session_state.pdf_text, question, max_chunks=7)
    previous = "\n".join(f"{item['role']}: {item['content']}" for item in conversation[-6:])
    prompt = f"""You are Studient AI Tutor. Teach using the workspace material as the primary source.
Student level: {level}
Tutor mode: {mode}
Student question: {question}
Recent conversation:
{previous}

Rules:
- In Socratic mode, do not immediately reveal the final answer. Ask one guiding question and give a small hint.
- In Hint mode, give a progressive hint, then ask the student to try again.
- In Explain mode, explain clearly at the student's level and offer a second analogy or representation.
- Detect likely misconceptions and label them gently as "Possible misconception" when relevant.
- If the workspace does not contain the answer, say so instead of inventing a document fact.
- End with one short check-for-understanding question.

WORKSPACE CONTEXT:
{context}"""
    return ask_groq(prompt, max_tokens=1800)


def generate_study_material(material_type, focus, difficulty):
    context = get_relevant_chunks(st.session_state.pdf_text, focus or "all important topics", max_chunks=8, max_context_chars=10000)
    instructions = {
        "Cheat sheet": "Create a concise exam-ready cheat sheet with headings, key facts, definitions, processes, formulas, and common mistakes.",
        "Formula sheet": "Extract and explain every important formula, define each variable, state units, and include when to use it.",
        "Glossary": "Create a glossary of the most important terms with clear one-to-two sentence definitions and examples.",
        "Concept comparison table": "Create a markdown comparison table of the most important related concepts, including similarities, differences, use cases, and exam traps.",
        "Timeline summary": "Create a chronological timeline with dates or sequence markers, events, causes, consequences, and memory cues.",
        "Case studies": "Create realistic case studies based only on the workspace, followed by analysis questions and model reasoning.",
        "Lab-viva questions": "Create laboratory viva questions with concise model answers, safety notes, observations, and common examiner follow-ups.",
        "Oral examination questions": "Create oral examination questions from easy to difficult, with ideal answer points and follow-up questions.",
    }
    prompt = f"""{instructions[material_type]}
Difficulty/depth: {difficulty}
Focus: {focus or 'all important topics'}
Return polished markdown suitable for a student to revise directly. Ground every factual claim in the workspace material and say when the material does not provide an answer.

WORKSPACE MATERIAL:
{context}"""
    return ask_groq(prompt, max_tokens=3600)


def generate_sample_paper(duration_minutes, total_marks, difficulty, section_mix, instructions):
    context = get_relevant_chunks(st.session_state.pdf_text, "all important exam questions concepts definitions processes formulas examples", max_chunks=18, max_context_chars=24000)
    prompt = f"""Create a complete university examination sample paper based strictly on the study material below.
Duration: {duration_minutes} minutes
Total marks: {total_marks}
Difficulty: {difficulty}
Requested sections: {section_mix}
Additional instructions: {instructions or 'Use a balanced university exam format.'}

Return polished markdown with:
1. University-style title, instructions, duration, total marks, and suggested time allocation.
2. Clearly numbered sections and questions with marks per question.
3. A balanced mix of important questions from across the entire material, not just one chapter.
4. Optional choices where appropriate.
5. A separate answer key/model-answer section after the paper with concise marking points and explanations.
6. A topic coverage note showing which major topics were tested.

Avoid invented facts. If the material lacks enough information for a question, omit it and use another supported concept.

FULL WORKSPACE STUDY MATERIAL:
{context}"""
    return ask_groq(prompt, max_tokens=6500)


def load_summary_history():
    return read_private_json(SUMMARY_HISTORY_FILE, [])


def save_summary(document_name, summary):
    if st.session_state.get("privacy_mode"):
        return
    history = load_summary_history()
    history.append({
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "document": document_name,
        "preview": clean_speech_text(summary)[:280],
    })
    try:
        write_private_json(SUMMARY_HISTORY_FILE, history[-20:])
    except Exception as e:
        st.warning(f"Couldn't save summary activity: {e}")


def load_review_state():
    return read_private_json(REVIEW_STATE_FILE, {})


def save_review_state(state):
    try:
        write_private_json(REVIEW_STATE_FILE, state)
    except Exception as e:
        st.warning(f"Couldn't save flashcard review progress: {e}")


def review_key(card):
    raw = f"{card.get('front', '')}|{card.get('topic', '')}|{st.session_state.get('document_name', '')}"
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]


def get_review_card(card):
    state = st.session_state.review_state.setdefault(review_key(card), {"box": 1, "due": str(date.today()), "history": [], "last_rating": "New"})
    return state


def rate_review_card(card, rating):
    state = get_review_card(card)
    intervals = {"Again": [0, 0, 1, 1, 1], "Hard": [1, 1, 2, 3, 4], "Good": [1, 2, 4, 7, 14], "Easy": [2, 4, 7, 14, 30]}
    if rating == "Again":
        state["box"] = 1
    elif rating == "Hard":
        state["box"] = max(1, state.get("box", 1))
    elif rating == "Good":
        state["box"] = min(5, state.get("box", 1) + 1)
    elif rating == "Easy":
        state["box"] = min(5, state.get("box", 1) + 2)
    days = intervals[rating][min(state["box"] - 1, 4)]
    state["due"] = str(date.today() + timedelta(days=days))
    state["last_rating"] = rating
    state.setdefault("history", []).append({"date": str(date.today()), "rating": rating, "box": state["box"]})
    state["history"] = state["history"][-30:]
    save_review_state(st.session_state.review_state)


def due_review_cards(cards):
    today = str(date.today())
    return [card for card in cards if get_review_card(card).get("due", today) <= today]


def dashboard_snapshot():
    attempts = load_history()
    summaries = load_summary_history()
    documents = {}
    for attempt in attempts:
        name = attempt.get("document", "Unknown document")
        documents[name] = attempt.get("timestamp", "")
    if st.session_state.get("document_name"):
        documents[st.session_state.document_name] = datetime.now().isoformat(timespec="seconds")

    total_questions = sum(int(item.get("total", 0)) for item in attempts)
    best_score = max((round(item["score"] / item["total"] * 100) for item in attempts if item.get("total")), default=0)
    current_flashcards = st.session_state.get("flashcards") or []
    mastered = sum(1 for card in current_flashcards if get_review_card(card).get("box", 1) >= 5)
    current_cards = len(current_flashcards)
    readiness = min(100, round(best_score * 0.7 + (20 if st.session_state.get("summary_text") else 0) + (10 if current_cards else 0)))

    days = sorted({item.get("timestamp", "")[:10] for item in attempts if item.get("timestamp")}, reverse=True)
    streak = 0
    if days:
        from datetime import date, timedelta
        today = date.today()
        for offset, day in enumerate(days):
            if day == str(today - timedelta(days=offset)):
                streak += 1
            else:
                break
    analytics = analytics_snapshot(st.session_state.get("document_name", "")) if st.session_state.get("document_name") else {"average": 0, "improvement": 0, "sessions": 0, "weakest": "Complete a practice test", "questions": 0}
    due_today = len(due_review_cards(current_flashcards)) if current_flashcards else 0
    recommended = f"Review {due_today} flashcards due today" if due_today else "Take a practice test" if total_questions == 0 else f"Review {analytics['weakest']} concepts"
    recent_docs = sorted(documents.items(), key=lambda item: item[1], reverse=True)[:5]
    recent_summaries = list(reversed(summaries[-4:]))
    return {"total_questions": total_questions, "best_score": best_score, "mastered": mastered, "current_cards": current_cards, "readiness": readiness, "streak": streak, "recommended": recommended, "recent_docs": recent_docs, "recent_summaries": recent_summaries, "average": analytics["average"], "improvement": analytics["improvement"], "sessions": analytics["sessions"], "weakest": analytics["weakest"], "due_today": due_today}


def render_dashboard():
    data = dashboard_snapshot()
    doc_rows = "".join(f'<div class="sm-dashboard-row"><span>📄 {_html_escape_lib.escape(name)}</span><span>{stamp[:10] or "Recently"}</span></div>' for name, stamp in data["recent_docs"]) or '<div class="sm-dashboard-empty">Upload a document to begin your study journey.</div>'
    summary_rows = "".join(f'<div class="sm-dashboard-summary"><b>{_html_escape_lib.escape(item.get("document", "Study summary"))}</b><span>{_html_escape_lib.escape(item.get("preview", ""))}</span></div>' for item in data["recent_summaries"]) or '<div class="sm-dashboard-empty">Your generated summaries will appear here.</div>'
    html(f'''<section class="sm-dashboard">
      <div class="sm-dashboard-heading"><div><div class="sm-progress-kicker">Personal learning command center</div><h2>Welcome back to your study space</h2><p>Track your momentum, pick up where you left off, and make every session count.</p></div><div class="sm-dashboard-recommend"><span>Next best step</span><b>✨ {data["recommended"]}</b></div></div>
      <div class="sm-dashboard-grid">
        <div class="sm-dashboard-gauge"><div class="sm-progress-orb" style="--progress:{data["readiness"]};"><div class="sm-progress-orb-content"><span class="sm-progress-number">{data["readiness"]}%</span><span class="sm-progress-label">readiness</span></div></div><div><div class="sm-dashboard-card-kicker">EXAM READINESS</div><b>Build confident recall</b><p>Your readiness grows as you summarize, practice, and review.</p></div></div>
        <div class="sm-dashboard-stat"><span>🔥</span><b>{data["streak"]} day</b><small>study streak</small></div>
        <div class="sm-dashboard-stat"><span>✅</span><b>{data["total_questions"]}</b><small>questions completed</small></div>
        <div class="sm-dashboard-stat"><span>🎴</span><b>{data["mastered"]}/{data["current_cards"]}</b><small>flashcards mastered</small></div>
        <div class="sm-dashboard-stat"><span>🏆</span><b>{data["best_score"]}%</b><small>best test score</small></div>
      </div>
      <div class="sm-dashboard-insights"><span>Average score <b>{data["average"]}%</b></span><span>Trend <b class="{'positive' if data['improvement'] >= 0 else 'negative'}">{data["improvement"]:+d}%</b></span><span>Study sessions <b>{data["sessions"]}</b></span><span>Due today <b>{data["due_today"]}</b></span><span>Focus next <b>{_html_escape_lib.escape(data["weakest"])}</b></span></div>
      <div class="sm-dashboard-columns"><div><div class="sm-dashboard-section-title">Recent documents</div>{doc_rows}</div><div><div class="sm-dashboard-section-title">Recently generated summaries</div>{summary_rows}</div></div>
    </section>''')
 
 
# ============================================================ 
# SESSION STATE 
# ============================================================ 
 
defaults = { 
    "pdf_text": "", "document_name": "", "document_type": "", "document_signature": "", 
    "quiz_questions": None, "quiz_index": 0, "quiz_answers": [], 
    "quiz_locked": False, "quiz_finished": False, "mind_map": None,
    "summary_text": "", "summary_audio": None,
    "theme_mode": "Light", "reduced_motion": False,
    "high_contrast": False, "font_scale": "Default",
    "flashcards": None, "flashcard_status": {},
    "workspace_documents": {}, "review_state": load_review_state(),
    "workspace_name": "My Study Workspace",
    "test_category": "Mixed",
    "timed_mode": False,
    "test_time_limit": 10,
    "negative_marking": False,
    "test_started_at": None,
    "test_timed_out": False,
    "incorrect_questions": [],
    "important_questions": None, "mcq_questions": None,
    "study_plan": None, "tutor_messages": [], "tutor_level": "Intermediate",
    "generated_material": None, "generated_material_type": "", "sample_paper": None,
    "privacy_mode": False, "retention_days": 0,
} 
for k, v in defaults.items(): 
    if k not in st.session_state: 
        st.session_state[k] = v 

if st.session_state.theme_mode == "Dark":
    html("""
    <style>
    :root { --ink:#eef2ff; --ink-soft:#c7d2fe; --ink-faint:#94a3b8; --glass-bg:rgba(30,41,75,.66); --glass-bg-soft:rgba(30,41,75,.52); --glass-border:rgba(148,163,184,.24); }
    .stApp { background:radial-gradient(circle at 8% 0%,rgba(79,70,229,.28),transparent 32%),radial-gradient(circle at 92% 4%,rgba(219,39,119,.20),transparent 30%),linear-gradient(160deg,#10152d,#17162f 55%,#21152f) !important; }
    section[data-testid="stSidebar"] { background:linear-gradient(180deg,rgba(15,23,42,.92),rgba(30,27,55,.9)) !important; }
    .sm-hero,.sm-card,.sm-feature,.sm-metric,.sm-document-bar { background:linear-gradient(145deg,rgba(30,41,75,.78),rgba(49,46,129,.42)) !important; border-color:rgba(148,163,184,.24) !important; }
    .stTabs [data-baseweb="tab-list"] { background:rgba(30,41,75,.68) !important; }
    </style>
    """)
if st.session_state.high_contrast:
    html("""
    <style>
    :root { --ink:#111827; --ink-soft:#1f2937; --ink-faint:#374151; --glass-border:rgba(17,24,39,.38); }
    .stApp { filter:saturate(1.18) contrast(1.06); }
    .sm-card,.sm-dashboard,.sm-feature,.sm-metric,.sm-document-bar,.sm-progress-dashboard { border-color:rgba(17,24,39,.32) !important; }
    .stCaption, [data-testid="stCaptionContainer"] { color:#374151 !important; }
    </style>
    """)
font_scale_values = {"Small": ".92", "Default": "1", "Large": "1.10", "Extra large": "1.22"}
html(f'<style>:root {{ --sm-font-scale:{font_scale_values.get(st.session_state.font_scale, "1")}; }} .stApp {{ font-size:calc(1rem * var(--sm-font-scale)); }} .sm-card, .sm-dashboard, .sm-feature, .sm-document-bar {{ font-size:calc(1em * var(--sm-font-scale)); }}</style>')
if st.session_state.reduced_motion:
    html("<style>* { animation-duration:0.001ms !important; transition-duration:0.001ms !important; }</style>")
 
 
def reset_quiz(): 
    st.session_state.quiz_questions = None 
    st.session_state.quiz_index = 0 
    st.session_state.quiz_answers = [] 
    st.session_state.quiz_locked = False 
    st.session_state.quiz_finished = False 
    st.session_state.test_started_at = None
    st.session_state.test_timed_out = False
 
 
def metric_card(chip_class, icon, label, value): 
    html(f""" 
    <div class="sm-metric"> 
      <div class="row"><span class="chip {chip_class}">{icon}</span><span class="label">{label}</span></div> 
      <div class="value">{value}</div> 
    </div> 
    """) 
 
 
# ============================================================ 
# HERO 
# ============================================================ 
 
html(""" 
<div class="sm-hero">
  <div class="eyebrow">✦ Your intelligent study companion</div>
  <div class="hero-icon">🧠</div>
  <h1><span class="grad">Studient</span> AI</h1>
  <p class="sub">Your Personal AI-Powered Study Assistant</p>
  <p class="desc">Turn your notes, documents, and presentations into clear, exam-ready learning — then practice, reflect, and improve with confidence.</p>
</div> 
""") 
 
# ============================================================ 
# SIDEBAR 
# ============================================================ 
 
with st.sidebar:
    html("""
    <div class="sm-brand">
      <div class="sm-brand-mark">🧠</div>
      <div><div class="sm-brand-name"><span>Studient</span> AI</div><span class="sm-brand-sub">Personal study assistant</span></div>
    </div>
    <div class="sm-upload-head">Build your study space</div>
    <div class="sm-upload-copy">Upload a source and let Studient AI turn it into active revision material.</div>
    <div class="sm-file-types"><span class="sm-file-type">PDF</span><span class="sm-file-type">WORD</span><span class="sm-file-type">PPTX</span><span class="sm-file-type">TXT</span><span class="sm-file-type">MD</span></div>
    """)
    selected_theme = st.selectbox("Appearance", ["Light", "Dark"], index=0 if st.session_state.theme_mode == "Light" else 1)
    reduced_motion = st.checkbox("Reduce animations", value=st.session_state.reduced_motion)
    high_contrast = st.checkbox("High contrast", value=st.session_state.high_contrast, help="Increase contrast for text, borders, and surfaces.")
    font_scale = st.selectbox("Text size", ["Small", "Default", "Large", "Extra large"], index=["Small", "Default", "Large", "Extra large"].index(st.session_state.font_scale))
    if selected_theme != st.session_state.theme_mode or reduced_motion != st.session_state.reduced_motion or high_contrast != st.session_state.high_contrast or font_scale != st.session_state.font_scale:
        st.session_state.theme_mode = selected_theme
        st.session_state.reduced_motion = reduced_motion
        st.session_state.high_contrast = high_contrast
        st.session_state.font_scale = font_scale
        st.rerun()
    workspace_name = st.text_input("Workspace name", value=st.session_state.get("workspace_name", "My Study Workspace"), label_visibility="collapsed", placeholder="Name this course workspace")
    with st.expander("🔐 Privacy & data controls"):
        privacy_mode = st.checkbox("Private session mode", value=st.session_state.privacy_mode, help="Do not save new test attempts or summaries to local history.")
        retention_options = {"Keep forever": 0, "Keep 7 days": 7, "Keep 30 days": 30, "Keep 90 days": 90}
        retention_label = st.selectbox("History retention", list(retention_options), index=list(retention_options.values()).index(st.session_state.retention_days))
        if privacy_mode != st.session_state.privacy_mode or retention_options[retention_label] != st.session_state.retention_days:
            st.session_state.privacy_mode = privacy_mode
            st.session_state.retention_days = retention_options[retention_label]
            apply_retention_policy(st.session_state.retention_days)
            st.rerun()
        export_data = make_user_data_export()
        st.download_button("⬇️ Export my study data", export_data, file_name="studient-data-export.zip", mime="application/zip", key="export_user_data")
        confirm_delete = st.checkbox("I understand this permanently deletes local study data", key="confirm_delete_data")
        if st.button("🗑️ Securely delete local data", key="delete_local_data", disabled=not confirm_delete):
            delete_local_study_data()
            st.success("Local study data was securely deleted from this app instance.")
            st.rerun()
        html('<div class="sm-privacy-note">API keys are read only from environment variables or Streamlit secrets and are never included in exports. This version has no account system; deletion removes local workspace history and session data.</div>')
    uploaded_files = st.file_uploader("Choose study materials", type=SUPPORTED_TYPES, accept_multiple_files=True, label_visibility="collapsed", help="Upload multiple PDF, DOCX, PPTX, TXT, or Markdown files")

    if uploaded_files:
        if sum(file.size for file in uploaded_files) > 50 * 1024 * 1024:
            st.error("This workspace is larger than 50 MB. Please upload a smaller set of documents.")
            st.stop()
        signature = "|".join(f"{file.name}:{file.size}" for file in uploaded_files)
        if st.session_state.document_signature != signature:
            documents = {}
            with st.spinner("Reading your workspace documents..."):
                progress = st.progress(0, text="Preparing workspace extraction...")
                for index, file in enumerate(uploaded_files, 1):
                    progress.progress((index - 1) / len(uploaded_files), text=f"Extracting {file.name} ({index}/{len(uploaded_files)})")
                    extracted, detected_type = extract_document_text(file)
                    documents[file.name] = {"type": detected_type, "text": clean_text(extracted), "size": file.size}
                progress.progress(1.0, text="Workspace extraction complete")
            combined = "\n\n".join(f"\n[DOCUMENT: {name}]\n{data['text']}" for name, data in documents.items() if data["text"])
            st.session_state.workspace_documents = documents
            st.session_state.pdf_text = combined
            st.session_state.workspace_name = workspace_name or "My Study Workspace"
            st.session_state.document_name = st.session_state.workspace_name
            st.session_state.document_type = "workspace"
            st.session_state.document_signature = signature
            st.session_state.summary_text = ""
            st.session_state.summary_audio = None
            st.session_state.flashcards = None
            st.session_state.flashcard_status = {}
            reset_quiz()
        if st.session_state.pdf_text:
            st.success(f"✓ Workspace ready · {len(st.session_state.workspace_documents)} document(s)")
            word_count = len(st.session_state.pdf_text.split())
            section_count = sum(st.session_state.pdf_text.count(marker) for marker in ["[PAGE ", "[OCR PAGE ", "[SLIDE ", "[DOCUMENT: ", "[TABLE "])
            st.metric("📖 Words", f"{word_count:,}")
            st.metric("🧩 Sections", section_count or "—")
            with st.expander("📚 Workspace files"):
                for name, data in st.session_state.workspace_documents.items():
                    st.write(f"**{name}** · {SUPPORTED_LABELS.get(data['type'], data['type'].upper())} · {len(data['text'].split()):,} words")

    st.divider()
    st.markdown("### ⚙️ Study Settings") 
    question_count = st.slider("Number of questions", 3, 15, 5) 
    difficulty = st.selectbox("Difficulty", ["Easy", "Medium", "Hard", "University Exam"]) 
    st.markdown("#### 🧪 Practice test options")
    test_category = st.selectbox("Test category", ["Mixed", "Definitions", "Processes", "Comparisons", "Facts", "Exam-focused"], key="test_category")
    timed_mode = st.checkbox("Timed mode", key="timed_mode")
    if timed_mode:
        st.selectbox("Time limit", [5, 10, 15, 20, 30], key="test_time_limit", format_func=lambda value: f"{value} minutes")
    st.checkbox("Negative marking", key="negative_marking", help="Wrong answers reduce the final score by 0.25 points.")
 
    st.divider() 
    if st.button("🔄 Reset Test"): 
        reset_quiz() 
        st.rerun() 
 
    st.divider() 
    html(""" 
    <div class="sm-card"> 
      <b>💡 Study Tip</b><br><br> 
      For best results, upload clear lecture notes, Word documents, or presentation slides. 
    </div> 
    """) 
 
# ============================================================ 
# WELCOME SCREEN 
# ============================================================ 
 
if not st.session_state.pdf_text: 
    html(""" 
    <div class="sm-card"> 
      <h2><span class="grad">👋 Welcome to Studient AI</span></h2> 
      <p>Upload your study material from the sidebar and transform it into personalized exam preparation.</p> 
    </div> 
    """) 
    c1, c2, c3 = st.columns(3) 
    for col, icon, title, text in [ 
        (c1, "📚", "Learn", "Generate summaries, explanations, definitions and key concepts."), 
        (c2, "📝", "Practice", "Generate MCQs, short questions, long questions and flashcards."), 
        (c3, "🎯", "Prepare", "Analyze difficulty and test yourself with a graded exam-style test."), 
    ]: 
        with col: 
            html(f'<div class="sm-feature"><div class="icon">{icon}</div><div class="title">{title}</div><div class="text">{text}</div></div>') 
    render_dashboard()
    html('<div class="sm-warning"><b>📌 OCR supported</b><br><br>Scanned PDFs are automatically processed with OCR when selectable text is not available.</div>') 
    st.stop() 
 
# ============================================================ 
# DOCUMENT HEADER  (Words / Characters / Difficulty — restored) 
# ============================================================ 
 
html(f'<div class="sm-document-bar"><div class="sm-document-name"><span>📄 {st.session_state.document_name}</span><span style="font-size:11px; padding:5px 9px; border-radius:999px; background:rgba(79,70,229,.09); color:var(--accent1); vertical-align:middle;">{SUPPORTED_LABELS.get(st.session_state.document_type, "DOCUMENT")}</span></div><div class="sm-document-meta">Ready for active learning</div></div>') 

with st.expander("📄 Preview extracted document text"):
    preview_text = st.session_state.pdf_text[:5000]
    st.text_area("Extracted content preview", preview_text, height=220, label_visibility="collapsed")
    st.download_button("⬇️ Download extracted text", st.session_state.pdf_text, file_name=f"{st.session_state.document_name.rsplit('.', 1)[0]}.txt", mime="text/plain", key="download_extracted_text")
 
m1, m2, m3 = st.columns(3) 
with m1: 
    metric_card("chip-blue", "📖", "Words", f"{len(st.session_state.pdf_text.split()):,}") 
with m2: 
    metric_card("chip-cyan", "🔤", "Characters", f"{len(st.session_state.pdf_text):,}") 
with m3: 
    metric_card("chip-orange", "🎯", "Difficulty", difficulty) 

render_dashboard()
 
st.markdown("<br>", unsafe_allow_html=True) 
 
# ============================================================ 
# FLAT TAB ROW — everything in one scrollable line 
# ============================================================ 
 
tabs = st.tabs([ 
    "📚 Summary", "📝 Questions", "❓ MCQs", "🎴 Flashcards", 
    "📖 Long Questions", "🎯 Short Questions", "🔍 Key Concepts", 
    "📊 Difficulty", "🧪 Practice Test", "📈 Analytics", "🧠 Mind Map", "🗓️ Study Plan", "🧑‍🏫 AI Tutor", "🧰 Study Materials", "📝 Sample Paper", "💬 Ask Workspace",
]) 
 
# ---------------- SUMMARY ---------------- 
with tabs[0]: 
    st.header("📚 Chapter Summary") 
    st.caption("Generate an organized, exam-focused summary.") 
    if st.button("✨ Generate Summary", key="gen_summary"): 
        context = get_relevant_chunks(st.session_state.pdf_text, "overview main ideas important concepts", max_chunks=10) 
        prompt = f"""Create a clear, organized, exam-focused summary of this material. 
Structure it with: 1. Chapter Overview 2. Main Concepts 3. Important Definitions 
4. Important Facts 5. Processes / Mechanisms 6. Exam-Focused Points 7. Quick Revision. 
 
STUDY MATERIAL: 
{context}""" 
        with st.spinner("Studient AI is analyzing your material..."): 
            result = ask_groq(prompt, max_tokens=2600) 
        if result: 
            st.session_state.summary_text = result
            st.session_state.summary_audio = None
            save_summary(st.session_state.document_name, result)
            st.rerun()
            html(f'<div class="sm-answer"><div class="title">📚 AI Study Summary</div></div>') 
            st.markdown(result) 
            st.download_button("⬇️ Download summary as Markdown", result, file_name="studient-summary.md", mime="text/markdown", key="download_summary_markdown")

    if st.session_state.get("summary_text"):
        html('<div class="sm-answer"><div class="title">📚 AI Study Summary</div></div>')
        st.markdown(st.session_state.summary_text)
        st.download_button("⬇️ Download summary as Markdown", st.session_state.summary_text, file_name="studient-summary.md", mime="text/markdown", key="download_summary_markdown_saved")
        summary_pdf = make_pdf_bytes("Studient AI Study Summary", [(st.session_state.document_name, st.session_state.summary_text)])
        if summary_pdf:
            st.download_button("⬇️ Download summary as PDF", summary_pdf, file_name="studient-summary.pdf", mime="application/pdf", key="download_summary_pdf")
        html('<div class="sm-audio-panel"><div class="sm-audio-title">🔊 Listen to your study summary</div><div class="sm-audio-copy">Choose a language and let Studient AI read the summary aloud while you revise.</div></div>')
        audio_col, button_col = st.columns([1, 1.25])
        with audio_col:
            speech_language = st.selectbox("Reading language", ["English", "Hindi", "Urdu", "Spanish", "French"], key="summary_speech_language")
        speech_codes = {"English": "en", "Hindi": "hi", "Urdu": "ur", "Spanish": "es", "French": "fr"}
        with button_col:
            st.markdown("<br>", unsafe_allow_html=True)
            if st.button("🎧 Generate audio reading", key="generate_summary_audio"):
                with st.spinner("Preparing your audio summary..."):
                    speech_text = prepare_speech_text(st.session_state.summary_text, speech_language)
                    st.session_state.summary_audio = generate_speech_audio(speech_text, speech_codes[speech_language])
        if st.session_state.get("summary_audio"):
            st.audio(st.session_state.summary_audio, format="audio/mp3")
            st.download_button("⬇️ Download audio summary", st.session_state.summary_audio, file_name="studient-summary.mp3", mime="audio/mpeg", key="download_summary_audio")
        study_pack = make_study_pack()
        st.download_button("📦 Download complete study pack", study_pack, file_name="studient-study-pack.zip", mime="application/zip", key="download_study_pack")
 
# ---------------- IMPORTANT QUESTIONS ---------------- 
with tabs[1]: 
    st.header("📝 Important Exam Questions") 
    st.caption("Generate questions based strictly on your uploaded material.") 
    if st.button("✨ Generate Important Questions", key="gen_iq"): 
        context = get_relevant_chunks(st.session_state.pdf_text, "important concepts topics definitions processes exam", max_chunks=8) 
        prompt = f"""Create {question_count} important university exam questions from the material below. 
Difficulty: {difficulty}. Prioritize core concepts, processes, definitions, comparisons, cause and effect. 
Return numbered questions only. 
 
STUDY MATERIAL: 
{context}""" 
        with st.spinner("Finding the most important questions..."): 
            result = ask_groq(prompt, max_tokens=1800) 
        if result: 
            lines = [l.strip() for l in result.split("\n") if l.strip()] 
            st.session_state.important_questions = lines[:question_count]
            for i, q in enumerate(lines[:question_count]): 
                html(f'<div class="sm-question"><div class="label">Question {i+1}</div></div>') 
                st.write(q) 
            questions_docx = make_questions_docx(st.session_state.important_questions, "Important Exam Questions")
            if questions_docx:
                st.download_button("⬇️ Export questions as DOCX", questions_docx, file_name="studient-important-questions.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", key="export_important_docx")
 
# ---------------- MCQS (interactive JSON, browsable) ---------------- 
with tabs[2]: 
    st.header("❓ Multiple Choice Questions") 
    st.caption("Generate exam-style MCQs with answers and explanations.") 
    if st.button("✨ Generate MCQs", key="gen_mcq"): 
        with st.spinner("Generating your MCQs..."): 
            questions = generate_quiz_questions(st.session_state.pdf_text, question_count, difficulty) 
        if questions: 
            st.session_state.mcq_questions = questions
            for i, q in enumerate(questions, 1): 
                with st.expander(f"Q{i}. {q['question']}"): 
                    for letter, opt in q["options"].items(): 
                        st.write(f"**{letter})** {opt}") 
                    st.success(f"Correct Answer: {q['correct']} — {q['explanation']}") 
            mcq_docx = make_questions_docx(questions, "Studient AI MCQs")
            if mcq_docx:
                st.download_button("⬇️ Export MCQs as DOCX", mcq_docx, file_name="studient-mcqs.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", key="export_mcq_docx")
 
# ---------------- FLASHCARDS ---------------- 
# ---------------- FLASHCARDS ----------------
with tabs[3]:
    st.header("🎴 Visual Flashcards")
    st.caption(
        "Click a card to flip it and reveal a visual explanation "
        "of the concept."
    )

    if st.button("✨ Generate Visual Flashcards", key="gen_flash"):

        with st.spinner("Creating your visual flashcards..."):

            flashcards = generate_flashcards(
                st.session_state.pdf_text,
                question_count,
                difficulty
            )

        if flashcards:
            st.session_state.flashcards = flashcards
            st.session_state.flashcard_status = {str(i): "New" for i in range(len(flashcards))}
            st.success(
                f"Created {len(flashcards)} visual flashcards!"
            )
            st.rerun()

        else:
            st.warning(
                "I couldn't create the flashcards. "
                "Please try generating them again."
            ) 

    if st.session_state.get("flashcards"):
        flashcards = st.session_state.flashcards
        render_flashcards(flashcards)
        due_cards = due_review_cards(flashcards)
        mastered_cards = sum(1 for card in flashcards if get_review_card(card).get("box", 1) >= 5)
        st.markdown(f"#### Daily review queue · **{len(due_cards)} due today**")
        st.caption(f"Mastered {mastered_cards}/{len(flashcards)} cards · Rate each card after reviewing it to schedule the next session.")
        rating_options = ["Again", "Hard", "Good", "Easy"]
        for i, card in enumerate(flashcards):
            state = get_review_card(card)
            label = f"Card {i + 1}: {card['front']}"
            with st.expander(f"{label} · {state.get('last_rating', 'New')} · due {state.get('due', str(date.today()))}"):
                rating = st.radio("How well did you recall it?", rating_options, horizontal=True, key=f"rating_{i}")
                if st.button("Save review rating", key=f"save_rating_{i}"):
                    rate_review_card(card, rating)
                    st.success(f"Saved {rating}. Next review: {get_review_card(card)['due']}")
                if state.get("history"):
                    st.caption("Learning history: " + " · ".join(f"{item['date']} {item['rating']}" for item in state["history"][-5:]))
        csv_buffer = StringIO()
        writer = csv.writer(csv_buffer)
        writer.writerow(["Front", "Topic", "Visual type", "Visual title", "Visual items", "Last rating", "Due date", "Leitner box"])
        for i, card in enumerate(flashcards):
            review = get_review_card(card)
            writer.writerow([card["front"], card["topic"], card["visual_type"], card["visual_title"], " | ".join(card["visual_items"]), review.get("last_rating", "New"), review.get("due", str(date.today())), review.get("box", 1)])
        st.download_button("⬇️ Export flashcards as CSV", csv_buffer.getvalue(), file_name="studient-flashcards.csv", mime="text/csv", key="export_flashcards_csv")
        flashcards_pdf = make_flashcards_pdf(flashcards)
        if flashcards_pdf:
            st.download_button("⬇️ Export flashcards as PDF", flashcards_pdf, file_name="studient-flashcards.pdf", mime="application/pdf", key="export_flashcards_pdf")
 
# ---------------- LONG QUESTIONS ---------------- 
with tabs[4]: 
    st.header("📖 Long Questions") 
    st.caption("Prepare detailed university examination questions.") 
    if st.button("✨ Generate Long Questions", key="gen_long"): 
        context = get_relevant_chunks(st.session_state.pdf_text, "process explain discuss compare analyze describe", max_chunks=8) 
        prompt = f"""Create {question_count} long-answer university examination questions. 
Difficulty: {difficulty}. Focus on Explain / Discuss / Analyze / Compare / Describe processes / Cause and effect. 
Return only numbered questions. 
 
STUDY MATERIAL: 
{context}""" 
        with st.spinner("Generating long questions..."): 
            result = ask_groq(prompt, max_tokens=1800) 
        if result: 
            html('<div class="sm-card"></div>') 
            st.markdown(result) 
            long_docx = make_questions_docx([line for line in result.splitlines() if line.strip()], "Long Examination Questions")
            if long_docx:
                st.download_button("⬇️ Export long questions as DOCX", long_docx, file_name="studient-long-questions.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", key="export_long_docx") 
 
# ---------------- SHORT QUESTIONS ---------------- 
with tabs[5]: 
    st.header("🎯 Short Questions") 
    st.caption("Practice definitions, facts and short conceptual questions.") 
    if st.button("✨ Generate Short Questions", key="gen_short"): 
        context = get_relevant_chunks(st.session_state.pdf_text, "definitions facts terms concepts differences", max_chunks=8) 
        prompt = f"""Create {question_count} short-answer questions from this study material. 
Focus on Definitions / Facts / Important terms / Differences / Basic concepts. 
Return only numbered questions. 
 
STUDY MATERIAL: 
{context}""" 
        with st.spinner("Generating short questions..."): 
            result = ask_groq(prompt, max_tokens=1800) 
        if result: 
            html('<div class="sm-card"></div>') 
            st.markdown(result) 
            short_docx = make_questions_docx([line for line in result.splitlines() if line.strip()], "Short Questions")
            if short_docx:
                st.download_button("⬇️ Export short questions as DOCX", short_docx, file_name="studient-short-questions.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", key="export_short_docx") 
 
# ---------------- KEY CONCEPTS ---------------- 
with tabs[6]: 
    st.header("🔍 Key Concepts") 
    st.caption("Extract the concepts you should know before your exam.") 
    if st.button("✨ Extract Key Concepts", key="gen_concepts"): 
        context = get_relevant_chunks(st.session_state.pdf_text, "main concepts important ideas definitions topics", max_chunks=8) 
        prompt = f"""Identify the most important concepts in the study material. 
For every concept provide: CONCEPT / Short explanation / Why it matters for the exam. 
 
STUDY MATERIAL: 
{context}""" 
        with st.spinner("Identifying key concepts..."): 
            result = ask_groq(prompt, max_tokens=2600) 
        if result: 
            html('<div class="sm-card"></div>') 
            st.markdown(result) 
 
# ---------------- DIFFICULTY ANALYSIS ---------------- 
with tabs[7]: 
    st.header("📊 Difficulty Analysis") 
    st.caption("Find the topics that may require the most preparation.") 
    if st.button("✨ Analyze Difficulty", key="gen_diff"): 
        context = get_relevant_chunks(st.session_state.pdf_text, "complex difficult advanced conceptual process", max_chunks=8) 
        prompt = f"""Analyze the difficulty of this study material. Provide: 
1. Overall difficulty 2. Easy topics 3. Medium topics 4. Difficult topics 
5. Concepts requiring memorization 6. Concepts requiring deep understanding 
7. Topics most likely to challenge students 8. Recommended study priority. 
 
STUDY MATERIAL: 
{context}""" 
        with st.spinner("Analyzing difficulty..."): 
            result = ask_groq(prompt, max_tokens=2600) 
        if result: 
            html('<div class="sm-card"></div>') 
            st.markdown(result) 
 
# ---------------- PRACTICE TEST (graded, interactive) ---------------- 
with tabs[8]: 
    st.header("🧪 Practice Test") 
    st.caption("Answers are hidden until you submit each question — no peeking.") 
 
    if st.session_state.quiz_questions is None: 
        if st.session_state.get("incorrect_questions"):
            if st.button("🔁 Retake Incorrect Questions", key="retake_incorrect"):
                st.session_state.quiz_questions = st.session_state.incorrect_questions
                st.session_state.quiz_index = 0
                st.session_state.quiz_answers = []
                st.session_state.quiz_locked = False
                st.session_state.quiz_finished = False
                st.session_state.test_started_at = time.time()
                st.session_state.test_timed_out = False
                st.rerun()
        if st.button("🚀 Start New Test", key="start_test"): 
            with st.spinner("Preparing your practice test..."): 
                questions = generate_quiz_questions(st.session_state.pdf_text, question_count, difficulty, st.session_state.test_category) 
            if questions: 
                st.session_state.quiz_questions = questions 
                st.session_state.quiz_index = 0 
                st.session_state.quiz_answers = [] 
                st.session_state.quiz_locked = False 
                st.session_state.quiz_finished = False 
                st.session_state.test_started_at = time.time()
                st.session_state.test_timed_out = False
                st.rerun() 
 
    elif not st.session_state.quiz_finished: 
        questions = st.session_state.quiz_questions 
        idx = st.session_state.quiz_index 
        q = questions[idx] 

        if st.session_state.timed_mode:
            if st_autorefresh is not None:
                st_autorefresh(interval=1000, limit=None, key="practice_timer_heartbeat")
            elapsed = int(time.time() - (st.session_state.test_started_at or time.time()))
            remaining = max(0, st.session_state.test_time_limit * 60 - elapsed)
            st.progress(remaining / (st.session_state.test_time_limit * 60), text=f"Time remaining: {remaining // 60:02d}:{remaining % 60:02d}")
            if remaining == 0:
                st.session_state.test_timed_out = True
                st.session_state.quiz_finished = True
                st.rerun()
 
        html(f'<div class="sm-quiz-progress">Question {idx + 1} of {len(questions)}</div>') 
        html(f'<div class="sm-quiz-topic">{q["topic"]}</div>') 
        html(f'<div class="sm-quiz-q">{q["question"]}</div>') 
 
        option_labels = [f"{letter}) {text}" for letter, text in q["options"].items()] 
        choice = st.radio("Choose an answer:", option_labels, key=f"choice_{idx}", 
                           disabled=st.session_state.quiz_locked, label_visibility="collapsed") 
        chosen_letter = choice.split(")")[0] if choice else None 
 
        if not st.session_state.quiz_locked: 
            if st.button("✅ Submit Answer", key=f"submit_{idx}"): 
                is_correct = chosen_letter == q["correct"] 
                st.session_state.quiz_answers.append({ 
                    "topic": q["topic"], "correct": is_correct, 
                    "chosen": chosen_letter, "answer": q["correct"], "question": q,
                }) 
                st.session_state.quiz_locked = True 
                st.rerun() 
        else: 
            last = st.session_state.quiz_answers[-1] 
            if last["correct"]: 
                st.markdown('<span class="sm-result-correct">✔ Correct!</span>', unsafe_allow_html=True) 
            else: 
                st.markdown(f'<span class="sm-result-wrong">✘ Not quite — correct answer: {q["correct"]}</span>', unsafe_allow_html=True) 
            st.caption(q["explanation"]) 
 
            if idx + 1 < len(questions): 
                if st.button("➡️ Next Question", key=f"next_{idx}"): 
                    st.session_state.quiz_index += 1 
                    st.session_state.quiz_locked = False 
                    st.rerun() 
            else: 
                if st.button("🏁 Finish Test", key="finish_test"): 
                    st.session_state.quiz_finished = True 
                    st.rerun() 
 
    else: 
        answers = st.session_state.quiz_answers 
        score = sum(1 for a in answers if a["correct"]) 
        total = len(answers) 
        penalty = sum(0.25 for a in answers if not a["correct"]) if st.session_state.negative_marking else 0
        final_points = max(0, score - penalty)
        pct = round((final_points / total) * 100) if total else 0
        st.session_state.incorrect_questions = [a.get("question") for a in answers if not a["correct"] and a.get("question")]
 
        topic_stats = defaultdict(lambda: {"correct": 0, "wrong": 0}) 
        for a in answers: 
            topic_stats[a["topic"]]["correct" if a["correct"] else "wrong"] += 1 
 
        if st.session_state.get("last_saved_index") != id(answers): 
            save_attempt(st.session_state.document_name, score, total, dict(topic_stats)) 
            st.session_state["last_saved_index"] = id(answers) 
            st.rerun()
 
        verdict = "Excellent! 🔥" if pct >= 80 else "Good work 👍" if pct >= 60 else "Keep practicing 💪"
        adaptive = "Try a harder difficulty next time." if pct >= 85 else "Review missed concepts before your next test." if pct < 60 else "Stay at this difficulty and build consistency."
        timing_note = " Time expired." if st.session_state.test_timed_out else ""
        html(f'<div class="sm-score-hero"><div class="big">{final_points:g}/{total}</div><div class="pct">{pct}% — {verdict}{timing_note}</div><div style="margin-top:8px;color:var(--ink-faint);font-size:13px;">{adaptive}</div></div>')
        if st.session_state.negative_marking:
            st.caption(f"Negative marking enabled: {penalty:.2f} point(s) deducted.")
        if st.session_state.incorrect_questions:
            st.markdown(f"**{len(st.session_state.incorrect_questions)} incorrect question(s)** are ready for a focused retake.")
            for i, answer in enumerate(answers, 1):
                if not answer["correct"]:
                    q = answer.get("question", {})
                    with st.expander(f"Review incorrect answer {i}: {q.get('question', 'Question')}"):
                        st.write(f"Correct answer: **{q.get('correct', answer.get('answer'))}**")
                        st.caption(q.get("explanation", "Review this concept in your study material."))
        test_sections = []
        for i, answer in enumerate(answers, 1):
            q = answer.get("question", {})
            test_sections.append((f"Question {i}: {q.get('question', '')}", f"Your answer: {answer.get('chosen', 'Not answered')}\nCorrect answer: {q.get('correct', answer.get('answer', ''))}\nExplanation: {q.get('explanation', '')}"))
        test_pdf = make_pdf_bytes(f"Studient AI Practice Test — {pct}%", test_sections)
        if test_pdf:
            st.download_button("⬇️ Export practice test as PDF", test_pdf, file_name="studient-practice-test.pdf", mime="application/pdf", key="export_practice_test_pdf")
 
        weak = [t for t, s in topic_stats.items() if s["wrong"] > s["correct"]] 
        strong = [t for t, s in topic_stats.items() if s["correct"] > s["wrong"]] 
 
        col1, col2 = st.columns(2) 
        with col1: 
            st.markdown("**Weak topics**") 
            html("".join(f'<span class="sm-tag-weak">{t}</span>' for t in weak) or "<i>None — nice.</i>") 
        with col2: 
            st.markdown("**Strong topics**") 
            html("".join(f'<span class="sm-tag-strong">{t}</span>' for t in strong) or "<i>None yet.</i>") 
 
        if st.button("🔄 Take Another Test", key="retake"): 
            reset_quiz() 
            st.rerun() 
 
# ---------------- ANALYTICS ---------------- 
with tabs[9]: 
    st.header("📈 Study Analytics") 
    history = load_history() 
    doc_history = [h for h in history if h["document"] == st.session_state.document_name] 
 
    data = analytics_snapshot(st.session_state.document_name)
    if not doc_history:
        html('<div class="sm-card"><p>No test attempts yet for this document. Take a Practice Test to see your analytics here.</p></div>')
    else:
        metric_cols = st.columns(5)
        metric_cols[0].metric("Average score", f"{data['average']}%")
        metric_cols[1].metric("Improvement", f"{data['improvement']:+d}%")
        metric_cols[2].metric("Questions", data["questions"])
        metric_cols[3].metric("Sessions", data["sessions"])
        metric_cols[4].metric("Study time", f"{sum(item.get('duration_minutes', 0) for item in doc_history):.1f} min")

        st.markdown("#### Score improvement over time")
        st.line_chart(data["scores"])
        st.caption("Accessible chart description: each point represents your percentage score for a completed practice test, ordered from oldest to newest.")
        if data["improvement"] > 0:
            st.success(f"Your score improved by {data['improvement']}% across your recorded attempts.")
        elif data["improvement"] < 0:
            st.info(f"Your score changed by {data['improvement']}%. Review the weakest concepts before retaking the test.")

        st.markdown("#### Topic-level accuracy")
        if data["topic_accuracy"]:
            st.bar_chart(data["topic_accuracy"])
            topic_text = " · ".join(f"{topic}: {accuracy}%" for topic, accuracy in sorted(data["topic_accuracy"].items(), key=lambda item: item[1]))
            st.caption(f"Accessible chart description: topic accuracy from lowest to highest. {topic_text}")

        st.markdown("#### Flashcard mastery")
        mastery_pct = round(data["mastered"] / data["cards"] * 100) if data["cards"] else 0
        st.progress(mastery_pct / 100, text=f"{data['mastered']}/{data['cards']} mastered · {mastery_pct}%")

        st.markdown("#### Recommended next topic")
        st.info(f"Review **{data['weakest']}** next, then practice the related flashcards before your next test.")
        st.caption("Analytics are stored locally to this app instance and reset if the app restarts or redeploys.")
 

# ---------------- INTERACTIVE MIND MAP ----------------
with tabs[10]:
    st.header("🧠 Interactive Mind Map")
    st.caption("Turn your uploaded material into a visual, collapsible concept hierarchy.")

    if st.button("✨ Generate Interactive Mind Map", key="gen_mind_map"):
        with st.spinner("Mapping the key ideas in your document..."):
            mind_map = generate_mind_map(st.session_state.pdf_text)
        if mind_map:
            st.session_state.mind_map = mind_map
            st.success("Your interactive mind map is ready.")
        else:
            st.warning("I couldn't create the mind map. Please try generating it again.")

    if st.session_state.get("mind_map"):
        render_mind_map(st.session_state.mind_map)
        st.download_button("⬇️ Export mind map as JSON", json.dumps(st.session_state.mind_map, indent=2, ensure_ascii=False), file_name="studient-mind-map.json", mime="application/json", key="export_mind_map_json")
        mindmap_png = make_mindmap_png(st.session_state.mind_map)
        if mindmap_png:
            st.download_button("⬇️ Export mind map as PNG", mindmap_png, file_name="studient-mind-map.png", mime="image/png", key="export_mind_map_png")
    else:
        html('<div class="sm-card"><p>Generate a mind map to explore your document visually. Each branch is grounded in the uploaded material.</p></div>')

# ---------------- ASK WORKSPACE ---------------- 
with tabs[15]: 
    st.header("💬 Ask Workspace / Outside Knowledge") 
    st.caption("Ask questions grounded in your uploaded workspace, or optionally let Studient AI add clearly labeled general knowledge.") 
 
    use_general_knowledge = st.checkbox( 
        "💡 Include outside / general knowledge", 
        value=False, 
        help="Off = strict, hallucination-safe answers grounded only in your PDF (best for studying facts). " 
             "On = the PDF is used as context, but the AI can add its own knowledge and ideas — " 
             "answers will clearly separate what came from your document vs. general knowledge.", 
    ) 
 
    user_question = st.text_input("Your question", placeholder="e.g. Explain the process of stellar evolution.") 
    if st.button("🤖 Ask Studient", key="ask_pdf"): 
        if not user_question.strip(): 
            st.warning("Please enter a question first.") 
        else: 
            context = get_relevant_chunks(st.session_state.pdf_text, user_question, max_chunks=7) 
 
            if use_general_knowledge: 
                prompt = f"""The student uploaded the document below and is asking a question that may go 
beyond what the document literally states (e.g. asking for improvements, opinions, or ideas). 
 
Answer helpfully using BOTH the document and your own general knowledge. Structure your answer 
in two clearly labeled parts: 
 
📄 From your document: 
(what the document itself says that's relevant — if nothing is relevant, say so briefly) 
 
💡 Suggestions / general knowledge: 
(your own ideas, recommendations, or knowledge that goes beyond the document — be specific and practical) 
 
Do not present your own knowledge as if it came from the document. 
 
DOCUMENT CONTENT: 
{context} 
 
STUDENT QUESTION: 
{user_question}""" 
            else: 
                prompt = f"""Answer the student's question using ONLY the information in the provided workspace material. 
If the answer cannot be found, say clearly: "The answer is not available in the uploaded PDF." 
Do not invent facts. Explain clearly, use bullet points if useful. 
 
WORKSPACE MATERIAL: 
{context} 
 
STUDENT QUESTION: 
{user_question}""" 
 
            with st.spinner("Searching your PDF and thinking..."): 
                answer = ask_groq(prompt, max_tokens=2500) 
            if answer: 
                title = "🤖 StudientAI Answer" if not use_general_knowledge else "🤖 StudientAI Answer (PDF + general knowledge)" 
                html(f'<div class="sm-answer"><div class="title">{title}</div></div>') 
                st.markdown(answer) 

# ---------------- STUDY PLAN ----------------
with tabs[11]:
    st.header("🗓️ Personalized Study Plan")
    st.caption("Turn your exam date, goals, availability, and weak topics into a practical daily plan grounded in your workspace.")
    plan_col1, plan_col2 = st.columns(2)
    with plan_col1:
        exam_date = st.date_input("Exam date", value=date.today() + timedelta(days=14), min_value=date.today(), key="plan_exam_date")
        target_grade = st.selectbox("Target grade", ["Pass", "B / Good", "A / Excellent", "A+ / Distinction"], key="plan_target_grade")
        hours_per_day = st.number_input("Available hours per day", min_value=0.5, max_value=12.0, value=2.0, step=0.5, key="plan_hours")
    with plan_col2:
        weak_subjects = st.text_area("Weak subjects or topics", placeholder="e.g. cellular respiration, processes, formulas", key="plan_weak_subjects")
        schedule = st.multiselect("Preferred study schedule", ["Morning", "Afternoon", "Evening", "Weekdays", "Weekends", "Short focused sessions"], default=["Evening", "Weekdays"], key="plan_schedule")
    if st.button("✨ Generate my study plan", key="generate_study_plan"):
        with st.spinner("Designing your personalized plan..."):
            st.session_state.study_plan = generate_study_plan(exam_date, target_grade, hours_per_day, weak_subjects, ", ".join(schedule))
        st.rerun()
    if st.session_state.get("study_plan"):
        html('<div class="sm-answer"><div class="title">🗓️ Your personalized study plan</div></div>')
        st.markdown(st.session_state.study_plan)
        st.download_button("⬇️ Download study plan", st.session_state.study_plan, file_name="studient-study-plan.md", mime="text/markdown", key="download_study_plan")

# ---------------- AI TUTOR ----------------
with tabs[12]:
    st.header("🧑‍🏫 AI Tutor")
    st.caption("Learn through guided questions, progressive hints, and explanations grounded in your workspace.")
    tutor_col1, tutor_col2 = st.columns(2)
    with tutor_col1:
        tutor_level = st.selectbox("Your level", ["Beginner", "Intermediate", "Advanced", "University exam"], key="tutor_level")
    with tutor_col2:
        tutor_mode = st.selectbox("Tutor mode", ["Socratic questions", "Progressive hints", "Explain with analogies"], key="tutor_mode")
    if st.session_state.tutor_messages:
        for message in st.session_state.tutor_messages:
            with st.chat_message(message["role"]):
                st.markdown(message["content"])
    tutor_question = st.text_area("Ask your tutor", placeholder="I don't understand how this process works...", key="tutor_question")
    if st.button("🧑‍🏫 Ask tutor", key="ask_tutor"):
        if not tutor_question.strip():
            st.warning("Ask a question so the tutor can guide you.")
        else:
            with st.spinner("Your tutor is thinking..."):
                reply = tutor_reply(tutor_question, tutor_level, tutor_mode, st.session_state.tutor_messages)
            if reply:
                st.session_state.tutor_messages.extend([{ "role": "user", "content": tutor_question }, { "role": "assistant", "content": reply }])
                st.rerun()
    if st.session_state.tutor_messages and st.button("🧹 Clear tutor conversation", key="clear_tutor"):
        st.session_state.tutor_messages = []
        st.rerun()

# ---------------- STUDY MATERIALS ----------------
with tabs[13]:
    st.header("🧰 Study Materials Lab")
    st.caption("Generate focused revision assets only when you need them, keeping large workspaces fast and uncluttered.")
    material_type = st.selectbox("Material type", ["Cheat sheet", "Formula sheet", "Glossary", "Concept comparison table", "Timeline summary", "Case studies", "Lab-viva questions", "Oral examination questions"], key="material_type")
    material_focus = st.text_input("Focus topic (optional)", placeholder="e.g. photosynthesis, Chapter 3, or all major concepts", key="material_focus")
    material_depth = st.selectbox("Depth", ["Quick revision", "Exam-ready", "Deep study"], key="material_depth")
    if st.button("✨ Generate study material", key="generate_study_material"):
        with st.spinner(f"Creating your {material_type.lower()}..."):
            st.session_state.generated_material = generate_study_material(material_type, material_focus, material_depth)
            st.session_state.generated_material_type = material_type
        st.rerun()
    if st.session_state.get("generated_material"):
        html(f'<div class="sm-answer"><div class="title">🧰 {st.session_state.generated_material_type}</div></div>')
        st.markdown(st.session_state.generated_material)
        st.download_button("⬇️ Download study material", st.session_state.generated_material, file_name=f"studient-{st.session_state.generated_material_type.lower().replace(' ', '-')}.md", mime="text/markdown", key="download_generated_material")

# ---------------- UNIVERSITY SAMPLE PAPER ----------------
with tabs[14]:
    st.header("📝 University Sample Paper")
    st.caption("Generate a complete exam-style paper from important questions across your entire study workspace, including a model-answer key.")
    paper_col1, paper_col2 = st.columns(2)
    with paper_col1:
        paper_duration = st.number_input("Exam duration (minutes)", min_value=30, max_value=300, value=120, step=15, key="paper_duration")
        paper_marks = st.number_input("Total marks", min_value=20, max_value=200, value=100, step=10, key="paper_marks")
        paper_difficulty = st.selectbox("Difficulty", ["Balanced university level", "Challenging", "Revision-friendly", "Final exam level"], key="paper_difficulty")
    with paper_col2:
        paper_sections = st.multiselect("Section mix", ["MCQs", "Short answers", "Long answers", "Problems / calculations", "Case analysis", "Essay questions"], default=["MCQs", "Short answers", "Long answers"], key="paper_sections")
        paper_instructions = st.text_area("Additional exam instructions", placeholder="e.g. Include questions from every chapter and emphasize processes.", key="paper_instructions")
    if st.button("✨ Generate full sample paper", key="generate_sample_paper"):
        with st.spinner("Building your university-style sample paper..."):
            st.session_state.sample_paper = generate_sample_paper(paper_duration, paper_marks, paper_difficulty, ", ".join(paper_sections), paper_instructions)
        st.rerun()
    if st.session_state.get("sample_paper"):
        html('<div class="sm-answer"><div class="title">📝 University sample paper with model answers</div></div>')
        st.markdown(st.session_state.sample_paper)
        st.download_button("⬇️ Download sample paper", st.session_state.sample_paper, file_name="studient-university-sample-paper.md", mime="text/markdown", key="download_sample_paper")

# ============================================================ 
# FOOTER 
# ============================================================ 
 
html(""" 
<div style="margin-top:44px; padding:26px; text-align:center; color:#94a3b8; font-size:13px; border-top:1px solid rgba(20,20,40,0.08);"> 
  🧠 StudientAI — All rights reserved to Sharjeelsarwar-ai2 
</div> 
""")
